"""Apply completed external evidence to both languages, preserving native math/Fig1."""
import csv,re,json,sys
from copy import deepcopy
from docx import Document
from docx.shared import Inches,Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from paths import *
from revision_text import RESULT_HEADING,RESULT_PARAS,METHODS,DISCUSSION,FIG_CAPTIONS,REFERENCE,ACK
MOTHER=PAPER/'交付稿_20260922_完整修订与中文对照'
def rows(name):
    with (RESULTS/name).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def n(v,d=2):return f'{float(v):,.{d}f}'
def put(p,s):
    p.clear()
    for i,t in enumerate(re.split(r'\^(\d+(?:,\d+)*)\^',s)):
        r=p.add_run(t)
        if i%2:r.font.superscript=True
    return p
def before(doc,anchor,s,style='Body Text'):
    p=doc.add_paragraph(style=style);put(p,s);anchor.addprevious(p._p);return p
def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr()
    for old in list(pr.findall(qn('w:shd'))):pr.remove(old)
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);pr.append(sh)
def format_table(t,cn=False,widths=None):
    t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    total=(8.5*2.54-4.6)/2.54
    if widths is None:
        values=[float(col.width.inches) for col in t.columns];widths=[x/sum(values)*total for x in values]
    else:widths=[x/sum(widths)*total for x in widths]
    pr=t._tbl.tblPr
    for tag in ['w:tblpPr','w:tblBorders','w:tblInd']:
        for e in list(pr.findall(qn(tag))):pr.remove(e)
    bd=OxmlElement('w:tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single' if side in ['top','bottom','insideH'] else 'nil');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D6D9DE');bd.append(el)
    pr.append(bd)
    grid=t._tbl.tblGrid
    for x in list(grid):grid.remove(x)
    for w in widths:
        x=OxmlElement('w:gridCol');x.set(qn('w:w'),str(Inches(w).twips));grid.append(x)
    tw=pr.find(qn('w:tblW'));tw.set(qn('w:w'),str(Inches(total).twips));tw.set(qn('w:type'),'dxa')
    for i,row in enumerate(t.rows):
        rp=row._tr.get_or_add_trPr()
        if rp.find(qn('w:cantSplit')) is None:rp.append(OxmlElement('w:cantSplit'))
        if i==0 and rp.find(qn('w:tblHeader')) is None:rp.append(OxmlElement('w:tblHeader'))
        emphasis=row.cells[0].text in ['R2 point','R2 点策略']
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);shade(cell,'F2F3F5' if i==0 else 'FFFFFF')
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_together=True
                p.paragraph_format.keep_with_next=(i<len(t.rows)-1) if len(t.rows)<=12 else i==0
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:
                    r.font.name='Times New Roman';r.font.size=Pt(9);r.font.bold=i==0 or emphasis;r.font.color.rgb=RGBColor(0,0,0)
                    rp2=r._r.get_or_add_rPr();fonts=rp2.find(qn('w:rFonts'))
                    if fonts is None:fonts=OxmlElement('w:rFonts');rp2.append(fonts)
                    fonts.set(qn('w:eastAsia'),'SimSun')
def add_table(doc,anchor,headers,data,widths,cn):
    t=doc.add_table(rows=1,cols=len(headers))
    for cell,v in zip(t.rows[0].cells,headers):cell.text=v
    for row in data:
        for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
    format_table(t,cn,widths);anchor.addprevious(t._tbl);return t
def style(doc,cn=False,supp=False,protected=None):
    for name,size in [('Normal',10.5 if supp else 11),('Body Text',10.5 if supp else 11),('Title',18),('Heading 1',12),('Heading 2',12),('Heading 3',11),('Caption',9)]:
        if name not in doc.styles:continue
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
        if name.startswith('Heading') or name=='Title':s.font.bold=True
        rp=s.element.get_or_add_rPr();ft=rp.find(qn('w:rFonts'))
        if ft is None:ft=OxmlElement('w:rFonts');rp.append(ft)
        ft.set(qn('w:eastAsia'),'Microsoft YaHei' if name.startswith('Heading') or name=='Title' else 'SimSun')
    for section in doc.sections:
        section.left_margin=section.right_margin=section.top_margin=section.bottom_margin=Cm(2.3)
    in_references=False
    for p in doc.paragraphs:
        if protected is not None and p._p is protected:continue
        pf=p.paragraph_format;pf.widow_control=True
        if p.style.name in ['Normal','Body Text']:
            pf.line_spacing=1.18;pf.space_after=Pt(5)
        if p.style.name.startswith('Heading'):
            pf.keep_with_next=True;pf.space_before=Pt(10);pf.space_after=Pt(5)
        if p._p.xpath('.//w:drawing'):
            pf.keep_with_next=True;pf.keep_together=True;pf.space_after=Pt(4);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if re.match(r'^(Fig\.\s*\d+\s*\||图\s*\d+\s*[｜|])',p.text):
            p.style=doc.styles['Caption'];pf.keep_together=True;pf.keep_with_next=False;pf.line_spacing=1.05;pf.space_after=Pt(8)
        if p.text in ['References','参考文献','Supplementary references','补充参考文献']:
            pf.page_break_before=False;in_references=True
        elif in_references:
            pf.line_spacing=1.0 if supp else Pt(11);pf.space_before=Pt(0);pf.space_after=Pt(1);pf.keep_together=True
            if not supp:
                pr=p._p.get_or_add_pPr();snap=pr.find(qn('w:snapToGrid'))
                if snap is None:snap=OxmlElement('w:snapToGrid');pr.append(snap)
                snap.set(qn('w:val'),'0')
            for r in p.runs:r.font.size=Pt(9.5)
        if p._p.xpath('.//m:oMathPara'):
            nxt=p._p.getnext()
            if nxt is not None:
                visible=''.join(nxt.xpath('.//w:t/text()'))
                if visible.startswith(('where','其中')):pf.keep_with_next=True
    for t in doc.tables:format_table(t,cn)
    doc.core_properties.title=doc.paragraphs[0].text;doc.core_properties.author='';doc.core_properties.last_modified_by=''
def caption(p,s):
    put(p,s)
    if ' | ' in s:label,tail=s.split(' | ',1);p.clear();r=p.add_run(label+' | ');r.bold=True;p.add_run(tail)
    elif '｜' in s:label,tail=s.split('｜',1);p.clear();r=p.add_run(label+'｜');r.bold=True;p.add_run(tail)
    p.style='Caption'

def article(lang):
    cn=lang=='CN';li=int(cn);source=MOTHER/f'NatureEnergy_ExecutableFlexibility_20260922_{lang}.docx';doc=Document(source);ps=list(doc.paragraphs)
    fig1=doc.inline_shapes[0];rid1=fig1._inline.graphic.graphicData.pic.blipFill.blip.embed;hash1=__import__('hashlib').sha256(doc.part.related_parts[rid1].blob).hexdigest();dim1=(fig1.width,fig1.height);caption1=ps[8].text
    abstract=ps[2].text+(' 匹配的 12 小时持续预约增量也在 Frontier 轨迹与第二市场年份中复现。' if cn else ' Matched 12-hour persistence gains also recur across Frontier traces and the secondary market year.')
    put(ps[2],abstract)
    if not cn:assert len(abstract.split())<=150
    put(ps[6],ps[6].text+(' 独立 Frontier 轨迹和 2024 年机制回放进一步检验这些执行结果。' if cn else ' Independent Frontier traces and a secondary 2024 market replay further test these execution mechanisms.'))
    put(ps[53],DISCUSSION[li])
    for number,idx in [(2,17),(3,31),(4,39),(5,46)]:caption(ps[idx],FIG_CAPTIONS[number][li])
    for shape,name in zip(list(doc.inline_shapes)[1:],['Fig2_service_floor','Fig3_economic_value','Fig4_risk_selection','Fig5_recourse']):
        p=FINAL/'02_figures'/f'{name}.png';rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed;doc.part.related_parts[rid]._blob=p.read_bytes()
        im=Image.open(p);shape.width=Cm(8.5*2.54-4.6);shape.height=int(shape.width*im.height/im.width)
    anchor=ps[47]._p
    before(doc,anchor,RESULT_HEADING[li],'Heading 2')
    before(doc,anchor,RESULT_PARAS[0][li])
    p=doc.add_paragraph();p.add_run().add_picture(str(FINAL/'02_figures/Fig6_external_validation.png'),width=Cm(8.5*2.54-4.6));anchor.addprevious(p._p)
    p=before(doc,anchor,'','Caption');caption(p,FIG_CAPTIONS[6][li])
    for pair in RESULT_PARAS[1:]:before(doc,anchor,pair[li])
    for heading_en,heading_cn,paras in METHODS:
        before(doc,ps[66]._p,heading_cn if cn else heading_en,'Heading 2')
        for pair in paras:before(doc,ps[66]._p,pair[li])
    before(doc,ps[112]._p,'致谢' if cn else 'Acknowledgements','Heading 2');before(doc,ps[112]._p,ACK[li])
    put(ps[113],ps[113].text+(' Frontier 的原始 SWF 日志与许可由 JSSPP 官方工作负载归档提供；配套外部验证归档保留原文件、字段头、映射系数、日历来源与结果表。' if cn else ' The official JSSPP archive provides the original Frontier SWF logs and licence; the external-validation archive retains the raw files, headers, mapping coefficients, calendar provenance and result tables.'))
    ref=doc.add_paragraph();ref._p.getparent().remove(ref._p);ref._p=deepcopy(ps[-1]._p);doc._element.body.insert(len(doc._element.body)-1,ref._p);put(ref,REFERENCE)
    style(doc,cn,False,ps[8]._p)
    # Preserve Figure 1's image bytes, size and exact caption content.
    assert __import__('hashlib').sha256(doc.part.related_parts[rid1].blob).hexdigest()==hash1 and (fig1.width,fig1.height)==dim1 and ps[8].text==caption1
    filename='NatureEnergy_ExecutableFlexibility_20260922_EXTVALID_FINAL'+('_CN' if cn else '')+'.docx'
    out=FINAL/'04_manuscript'/filename;doc.save(out)
    mi=next(i for i,p in enumerate(doc.paragraphs) if p.text in ['Methods','方法'])
    mainwords=sum(len(p.text.split()) for p in doc.paragraphs[3:mi] if not p.text.startswith(('Fig.','图','Table','表')))
    if not cn:assert mainwords<=3000
    return dict(file=str(out),paragraphs=len(doc.paragraphs),tables=len(doc.tables),figures=len(doc.inline_shapes),native_equations=len(doc.element.xpath('.//m:oMath')),figure1_sha256=hash1,figure1_caption_unchanged=True,abstract_words=len(abstract.split()) if not cn else None,main_words=mainwords if not cn else None)

def supplement(lang):
    cn=lang=='CN';li=int(cn);choose=lambda en,zh:zh if cn else en
    doc=Document(MOTHER/f'Supplementary_ExecutableFlexibility_20260922_{lang}.docx')
    anchor=next(p._p for p in doc.paragraphs if p.text in ['Supplementary references','补充参考文献'])
    def para(s,heading=False):return before(doc,anchor,s,'Heading 2' if heading else 'Body Text')
    def note(number,title_en,title_cn,body_en,body_cn):para(choose(f'Supplementary Note {number} | {title_en}',f'补充说明 {number}｜{title_cn}'),True);para(choose(body_en,body_cn))
    def table(number,title_en,title_cn,cap_en,cap_cn,heads_en,heads_cn,data,widths):
        para(choose(f'Supplementary Table {number} | {title_en}',f'补充表 {number}｜{title_cn}'),True)
        p=para(choose(cap_en,cap_cn));p.paragraph_format.keep_with_next=True
        add_table(doc,anchor,heads_cn if cn else heads_en,data,widths,cn)
    note(19,'Frontier fields and normalisation','Frontier 字段与归一化',METHODS[0][2][0][0].replace('^16^',''),METHODS[0][2][0][1].replace('^16^',''))
    para(METHODS[0][2][1][li])
    note(20,'Secondary market year and physical inputs','第二市场年份与物理输入',METHODS[1][2][0][0],METHODS[1][2][0][1]);para(METHODS[1][2][1][li])
    note(21,'Cohort conservation and causal execution','批次守恒与因果执行',
        'The original reservation, myopic-hedge and deadline-priority kernels are reused with the target calendar length. They admit a cohort only at its release and retain late work in later slots. Cohort-assignment exports reproduce both released totals and the final capacity calendar. For recourse, execution records appended at their historical execution times reproduce the final in-year book with zero maximum discrepancy; subsequent repair therefore leaves completed history unchanged. The 2025 Eagle 12 h costs and savings reconcile to their archived values within US$0.01, and source-file hashes remain unchanged. Daily CVaR and drawdown use the same mark-to-market service-adjusted account as the main paper.',
        '复用原预约、短视可行性保护和截止时间优先内核，仅传入目标日历长度。任务批次只在释放时进入，逾期工作保留于后续时隙。导出的批次分配同时还原释放总量和最终容量日历。重排在历史实际执行时刻追加的记录与最终年度日历完全一致，最大偏差为零，因此后续修复不改变已完成历史。2025 年 Eagle 的 12 小时成本与节省在 US$0.01 以内还原归档值，源文件哈希保持不变。日 CVaR 和回撤采用与主文相同的逐期服务调整口径。')
    note(22,'What reproduces and what does not','复现与未复现的结果',
        'All four primary 12 h pairs have positive Published-DA savings, a positive persistent-minus-myopic increment, and no late work with the hedge. Recourse removes every fixed-reservation failure that occurs in the primary mixed tests; Frontier-2024 already has zero late work, so that case does not demonstrate an active repair. The unhedged control misses work in every primary 12 h case. Universal persistence superiority is not reproduced: the 2025 24 h increments are −US$1,827.61 for Eagle and −US$2,306.56 for Frontier. At 7 h on Eagle-2025, the physical floor is satisfied but price-responsive online allocation still misses work. A feasible offline/earliest-service floor is not a guarantee for every causal market policy. No window, scale, status or day was removed in response to these outcomes.',
        '四组主 12 小时配对均具有正的已公布 DA 节省和持续预约减短视的正增量，带保护策略均无逾期。主混合测试中实际出现的全部固定预约逾期均被重排消除；Frontier-2024 本来就无逾期，因此该案例不证明主动修复作用。四组主 12 小时案例的无保护控制均有工作逾期。持续预约普遍占优没有复现：2025 年 24 小时增量在 Eagle 为 −US$1,827.61，在 Frontier 为 −US$2,306.56。Eagle-2025 的 7 小时窗口满足物理下限，但价格响应的在线分配仍出现逾期。离线或尽早执行可行的物理下限，并不保证每一种因果市场策略都可行。没有依据这些结果删除窗口、缩放、任务状态或日期。')
    note(23,'Trace and external-validity boundaries','轨迹与外部有效性边界',
        'The source header calls job status always 1, but the released rows also contain 0, −1 and −2. All are retained as submitted demand; job outcomes are not a selection criterion. There are 338 epoch-zero jobs in Frontier-2024 and 1,000 in Frontier-2025. The latter origin burst determines the primary 3 h floor; no warm-up deletion is used. The archive does not establish observed NYISO co-location or commercial SLAs. Runtime is retained only for characterisation. The resource-field sensitivity preserves the direction of the 12 h persistence effect while materially changing floors and savings, showing that the exact numerical floor depends on the demand proxy. These are two trace years from one additional system, not two independent facility deployments. The Frontier source acknowledgement and licence are included with the raw data.',
        '源日志头称任务状态总为 1，但发布记录还含 0、−1、−2。全部保留为提交需求，任务结果不作为筛选依据。Frontier-2024 有 338 项零时刻任务，Frontier-2025 有 1,000 项；后者的起点突发决定主代理的 3 小时下限，没有删除预热期。归档不能证明 NYISO 实际共址或商业 SLA。运行时长仅用于特征描述。资源字段敏感性保留 12 小时持续预约作用的方向，但显著改变下限和节省，表明具体数值依赖需求代理。这是一个额外系统的两个轨迹年份，不是两个独立设施部署。原始数据附带 Frontier 致谢与许可。')
    chars=rows('frontier_trace_characteristics.csv')
    table(31,'Independent Frontier trace characteristics','独立 Frontier 轨迹特征',
        'Node counts use the explicit header-defined resource property. Interarrival quantities are seconds; daily CV and zero-arrival fractions include every calendar interval. Full runtime and raw-resource distributions are retained in frontier_trace_characteristics.csv.',
        '节点数采用日志头明确定义的资源属性。到达间隔单位为秒，日 CV 和零到达比例包含全部日历区间。完整运行时长与原资源分布保留在对应 CSV。',
        ['Trace','Jobs','Days','Mean gap s','P95 gap s','Daily CV','Zero bins %'],['轨迹','任务数','天数','平均间隔 秒','P95间隔 秒','日CV','零到达 %'],
        [[r['trace'],n(r['number_of_jobs'],0),n(r['calendar_days'],0),n(r['mean_interarrival_seconds']),n(r['p95_interarrival_seconds']),n(r['daily_workload_cv'],3),n(100*float(r['zero_arrival_interval_fraction']))] for r in chars],[1.1,1.1,.6,1,1,.7,1])
    floors=rows('trace_service_floor.csv')
    table(32,'Physical floors across primary trace year pairs','主轨迹年份配对的物理下限',
        'Energy is annual MWh-IT; floors are hours. Frontier years share the exact Eagle-2025 annual normalisation. Eagle-2024 retains the registered map and hence a 0.84% larger annual service total.',
        '电量为全年 MWh-IT，下限为小时。Frontier 两年均使用 Eagle-2025 精确全年归一化。Eagle-2024 保留已登记映射，因此全年服务量高 0.84%。',
        ['Trace','Year','Annual IT','Peak prefix units','100% floor','75% floor','50% floor'],['轨迹','年份','全年IT','最大前缀工作单位','100%下限','75%下限','50%下限'],
        [[r['trace'],r['market_year'],n(r['annual_variable_it_mwh'],3),n(r['peak_outstanding_prefix'],3),n(r['physical_floor_h_100pct']),n(r['physical_floor_h_75pct']),n(r['physical_floor_h_50pct'])] for r in floors],[1,.65,1.1,1.15,.9,.9,.9])
    allrows=rows('all_mechanism_runs.csv');primary=[r for r in allrows if r['proxy']!='field7_sensitivity' and r['policy'] in ['earliest','myopic_da','persistent_da']]
    pnames={'earliest':('Earliest','尽早'),'myopic_da':('Myopic DA','短视DA'),'persistent_da':('Persistent DA','持续DA'),'unhedged_myopic_da':('Unhedged DA','无保护DA')}
    table(33,'Complete primary mechanism outcomes','完整主机制结果',
        'Costs and savings are US dollars; cost is service-adjusted and includes the common DA position. Late work is MWh-IT. All prespecified 7/12/24 h windows remain, including 7 h failures and 24 h persistence reversals.',
        '成本与节省为美元，成本为包含共同日前头寸的服务调整口径；逾期为 MWh-IT。保留全部预先指定的 7/12/24 小时窗口，包括 7 小时失败和 24 小时持续预约反转。',
        ['Trace year','h','Policy','Adjusted cost','Saving','Saving %','Late IT'],['轨迹 年份','h','策略','调整后成本','节省','节省 %','逾期IT'],
        [[r['trace']+' '+r['market_year'],r['service_window_h'],pnames[r['policy']][li],n(r['annual_adjusted_cost_usd']),n(r['saving_usd']),n(r['saving_rate_pct']),n(r['late_it_mwh'],3)] for r in primary],[1.05,.35,1.05,1.2,1.15,.7,1])
    summary=rows('external_validity_summary.csv')
    table(34,'Matched persistence contribution at 12 h','12 小时匹配持续预约贡献',
        'Saving is relative to each trace/year’s earliest-service adjusted cost. Monetary values are US dollars. Persistence share divides the increment by persistent-policy saving, not by total facility cost.',
        '节省相对于各轨迹年份自身的尽早执行调整后成本，金额为美元。持续预约占比以增量除以持续策略节省额，不以设施总成本为分母。',
        ['Trace year','Earliest cost','Myopic saving','Persistent saving','Increment','Share %'],['轨迹 年份','尽早成本','短视节省','持续节省','增量','占比 %'],
        [[r['trace']+' '+r['market_year'],n(r['earliest_cost']),n(r['myopic_da_saving']),n(r['persistent_da_saving']),n(r['persistence_increment']),n(r['persistence_share_pct'])] for r in summary],[1.2,1.15,1.15,1.15,1.1,.75])
    mixed=[r for r in rows('all_recourse_runs.csv') if r['proxy']!='field7_sensitivity']
    table(35,'Mixed contract recourse across traces and years','跨轨迹年份的混合合同重排',
        'The mix is 25/50/25% at 4/12/24 h. Late work is MWh-IT and values are US dollars. Frontier-2024 has no fixed-reservation failure; its zero-to-zero row is preserved.',
        '混合份额为 4/12/24 小时的 25/50/25%。逾期为 MWh-IT，价值为美元。Frontier-2024 没有固定预约失败，保留其零至零结果。',
        ['Trace year','Fixed late','Recourse late','Fixed value','Recourse value','Value change'],['轨迹 年份','固定逾期','重排逾期','固定价值','重排价值','价值变化'],
        [[r['trace']+' '+r['market_year'],n(r['fixed_late_mwh'],3),n(r['recourse_late_mwh'],3),n(r['fixed_value_usd']),n(r['recourse_value_usd']),n(r['value_change_usd'])] for r in mixed],[1.2,.9,.9,1.25,1.25,1])
    f_all=rows('trace_service_floor_all_proxies.csv');m_all=rows('all_recourse_runs.csv');sensitivity=[]
    for year in ['2024','2025']:
        for proxy in ['requested_nodes','field7_sensitivity']:
            g=[r for r in allrows if r['trace']=='Frontier' and r['market_year']==year and r['proxy']==proxy and r['service_window_h']=='12']
            persistent=next(r for r in g if r['policy']=='persistent_da');myopic=next(r for r in g if r['policy']=='myopic_da')
            f=next(r for r in f_all if r['trace']=='Frontier' and r['market_year']==year and r['proxy']==proxy)
            m=next(r for r in m_all if r['trace']=='Frontier' and r['market_year']==year and r['proxy']==proxy)
            sensitivity.append([year,choose('Nodes','节点') if proxy=='requested_nodes' else choose('Field 7','字段7'),n(f['physical_floor_h_100pct']),n(persistent['saving_usd']),n(float(persistent['saving_usd'])-float(myopic['saving_usd'])),n(m['fixed_late_mwh'],3),n(m['recourse_late_mwh'],3)])
    table(36,'Frontier resource proxy sensitivity','Frontier 资源代理敏感性',
        'Each proxy is separately normalised to 2,981.607392515 MWh-IT before evaluation. Savings and increments use 12 h; late work is from the mixed contract. No proxy is selected by performance.',
        '每个代理在评估前独立归一化为 2,981.607392515 MWh-IT。节省和增量使用 12 小时，逾期来自混合合同。不依据表现选择代理。',
        ['Year','Proxy','Floor h','Persistent saving','Increment','Fixed late','Recourse late'],['年份','代理','下限h','持续节省','增量','固定逾期','重排逾期'],sensitivity,[.65,.9,.7,1.4,1.1,.9,1])
    diag=rows('unhedged_execution_diagnostic.csv')
    table(37,'Unhedged myopic execution diagnostic','无保护短视执行诊断',
        'The 12 h price-argmin control removes only the deadline-prefix hedge. It retains all late work. Positive saving cannot establish service feasibility. Values and CVaR are US dollars; energy is MWh-IT.',
        '12 小时最低价控制仅移除截止日前缀保护，全部逾期工作继续保留。正节省不能证明服务可行。价值和 CVaR 为美元，电量为 MWh-IT。',
        ['Trace year','Saving','Shifted IT','Late IT','CVaR95'],['轨迹 年份','节省','调节IT','逾期IT','CVaR95'],
        [[r['trace']+' '+r['market_year'],n(r['saving_usd']),n(r['shifted_it_mwh'],3),n(r['late_it_mwh'],3),n(r['cvar95_usd'])] for r in diag],[1.5,1.3,1.3,1.3,1.2])
    tail=rows('nyiso_2024_observed_price_tail.csv')
    table(38,'Observed price settlement tail for 2024','2024 实测价格结算尾段',
        'The 48 h tail uses observed January 2025 RT. All evaluated 12/24 h point policies already close with no unfinished work, so the incremental tail effect is zero. This is an empty-obligation boundary check, not evidence from a new operating year.',
        '48 小时尾段使用 2025 年 1 月实测 RT。受测 12/24 小时点策略均已在年末清空工作，因此增量尾段效应为零。这是空义务边界检查，不是新增运行年份的证据。',
        ['Trace','Window h','Policy','Closing IT','Value change'],['轨迹','窗口h','策略','期末IT','价值变化'],
        [[r['trace'],r['service_window_h'],pnames[r['policy']][li],n(r['closing_work_mwh'],3),n(r['difference_usd'])] for r in tail],[1.1,1,1.7,1.3,1.3])
    # The full acknowledgement is in the article and the raw licences. Keep the
    # supplemental citation in its bibliography without a duplicate end block.
    ref=doc.add_paragraph();ref._p.getparent().remove(ref._p);ref._p=deepcopy(doc.paragraphs[-1]._p);doc._element.body.insert(len(doc._element.body)-1,ref._p);put(ref,REFERENCE)
    style(doc,cn,True)
    out=FINAL/'05_supplementary'/('Supplementary_ExecutableFlexibility_20260922_EXTVALID_FINAL'+('_CN' if cn else '')+'.docx');doc.save(out)
    return dict(file=str(out),tables=len(doc.tables),paragraphs=len(doc.paragraphs),native_equations=len(doc.element.xpath('.//m:oMath')))

if __name__=='__main__':
    result=[article('EN'),article('CN'),supplement('EN'),supplement('CN')]
    save(REPORTS/'document_build_manifest.json',result)
    print(json.dumps(result,indent=2,ensure_ascii=False))
