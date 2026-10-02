"""Apply the DART-ValueSelect protocol amendment to the revised manuscript.

This script performs a structural DOCX edit only.  It intentionally does not
render the document or carry out visual/layout review.
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


MARKER = "5.9 Decision-oriented evaluation and model selection for DART forecasting"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def package_snapshot(path: Path) -> dict[str, int]:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        archive.testzip()
    doc = Document(path)
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "inline_shapes": len(doc.inline_shapes),
        "styles": len(doc.styles),
        "package_parts": len(names),
        "media_parts": sum(name.startswith("word/media/") for name in names),
    }


def paragraph_with_text(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise KeyError(f"Paragraph not found: {text}")


def paragraph_after_heading(doc: Document, heading: str) -> Paragraph:
    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs[:-1]):
        if paragraph.text.strip() == heading:
            return paragraphs[index + 1]
    raise KeyError(f"Heading not found: {heading}")


def set_text(paragraph: Paragraph, text: str, changes: list[str], label: str) -> None:
    paragraph.text = text
    changes.append(label)


def insert_paragraph_after(
    anchor: Paragraph, text: str, style: str | None = None
) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def find_table(doc: Document, first_header: str, second_header: str | None = None):
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers and headers[0] == first_header and (
            second_header is None or (len(headers) > 1 and headers[1] == second_header)
        ):
            return table
    raise KeyError(f"Table not found: {first_header} / {second_header}")


def append_row(table, values: list[str]) -> None:
    row = table.add_row()
    for cell, value in zip(row.cells, values, strict=True):
        cell.text = value


def insert_rows_before_label(table, before_label: str, rows: list[list[str]]) -> None:
    target = next(row for row in table.rows if row.cells[0].text.strip() == before_label)
    for values in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
        target._tr.addprevious(row._tr)


def add_template_table_after(doc: Document, anchor, title: str, headers, rows, style_name):
    caption = doc.add_paragraph(title)
    caption.style = "Caption"
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = style_name
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for values in rows:
        append_row(table, list(values))
    anchor.addnext(caption._p)
    caption._p.addnext(table._tbl)
    return table._tbl


def main() -> None:
    target = Path(os.environ["DART_DOCX_PATH"])
    backup = Path(
        os.environ.get(
            "DART_DOCX_BACKUP",
            str(target.with_name(target.stem + "_pre_DART_ValueSelect_20260810.docx")),
        )
    )
    manifest_path = Path(os.environ["DART_DOCX_MANIFEST"])
    manifest_path.parent.mkdir(parents=True, exist_ok=True)

    original_hash = sha256(target)
    before = package_snapshot(target)
    original_doc = Document(target)
    already_applied = any(p.text.strip() == MARKER for p in original_doc.paragraphs)

    if already_applied:
        alignment_changes: list[str] = []
        for table in original_doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if "dart_probability_actionability.csv" in cell.text:
                        cell.text = cell.text.replace(
                            "dart_probability_actionability.csv",
                            "dart_probability_quality.csv",
                        )
                        alignment_changes.append(
                            "align_probability_table_artifact_name_with_specification"
                        )
                    if "high-RT-spread/high-workload/low-PV" in cell.text:
                        cell.text = cell.text.replace(
                            "high-RT-spread/high-workload/low-PV",
                            "large positive/negative canonical DART, high-workload, and low-PV",
                        )
                        alignment_changes.append(
                            "replace_ambiguous_rt_spread_extreme_label"
                        )
                    if "RT-spread diagnostic" in cell.text:
                        cell.text = cell.text.replace(
                            "RT-spread diagnostic",
                            "canonical DA-minus-RT DART metrics and explicitly named inverse RT-minus-DA diagnostic",
                        )
                        alignment_changes.append(
                            "replace_ambiguous_rt_spread_metric_label"
                        )
        if alignment_changes:
            original_doc.save(target)
        existing = {}
        if manifest_path.is_file():
            existing = json.loads(manifest_path.read_text(encoding="utf-8"))
        prior_alignment_changes = set(existing.get("post_patch_alignment_changes", []))
        if any(
            "dart_probability_quality.csv" in cell.text
            for table in original_doc.tables
            for row in table.rows
            for cell in row.cells
        ):
            prior_alignment_changes.add(
                "align_probability_table_artifact_name_with_specification"
            )
        manifest = {
            **existing,
            "status": "APPLIED",
            "target": str(target),
            "output_sha256": sha256(target),
            "structure_after": package_snapshot(target),
            "post_patch_alignment_changes": sorted(
                prior_alignment_changes.union(alignment_changes)
            ),
            "render_or_visual_review_performed": False,
        }
        manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
        print(json.dumps(manifest, indent=2))
        return

    if backup.exists():
        if sha256(backup) != original_hash:
            raise RuntimeError(f"Existing backup does not match source: {backup}")
    else:
        shutil.copy2(target, backup)

    doc = original_doc
    changes: list[str] = []

    # Abstract, positioning, and contribution boundary.
    abstract = paragraph_after_heading(doc, "Abstract")
    set_text(
        abstract,
        abstract.text.rstrip()
        + " DART is defined canonically as day-ahead minus real-time LBMP, and candidate probabilistic DART models are evaluated through DART-ValueSelect: statistical credibility and actionability are admission diagnostics, while the final ranking uses the lower confidence bound of counterfactual risk-adjusted operational value under an identical constrained optimizer.",
        changes,
        "abstract_dart_value_select",
    )
    core = paragraph_with_text(
        doc,
        "This question differs from ordinary data-center electricity-cost minimization in four ways. First, the flexible resource is a collection of tasks with engineered service and mobility constraints. Second, an operating envelope identified from referenced NLR source records is used to calibrate and test the physical twin before market optimization. Third, the NYISO case is a two-settlement counterfactual in which independently forecast day-ahead and real-time prices expose a frozen twin to market uncertainty. Fourth, the objective includes service quality, downside risk, transfer uncertainty, and grid-facing proxies rather than only expected energy cost.",
    )
    set_text(
        core,
        core.text.rstrip()
        + " The derived DART distribution is therefore judged by the decisions it supports for this constrained user, not by unconditional point error alone.",
        changes,
        "core_question_dart_boundary",
    )
    contribution_anchor = paragraph_with_text(
        doc,
        "A two-stage stochastic MILP with CVaR that evaluates counterfactual day-ahead procurement, real-time settlement, workload placement, migration, storage, PV, service, and grid-facing outcomes in NYISO.",
    )
    insert_paragraph_after(
        contribution_anchor,
        "DART-ValueSelect, which compares candidate probabilistic DART models in an identical rolling stochastic-optimization and realized-replay pipeline, filters infeasible or excessive-risk candidates, and selects only from a Model Confidence Set using the lower confidence bound of incremental operational value.",
        "List Number",
    )
    changes.append("contribution_dart_value_select")
    organization = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("Section 2 reviews")))
    set_text(
        organization,
        organization.text.replace(
            "Section 5 presents local physical calibration and domain-specific forecasting;",
            "Section 5 presents local physical calibration, domain-specific forecasting, and decision-oriented DART evaluation;",
        ).replace(
            "Section 8 preregisters the experimental protocol.",
            "Section 8 preregisters the experimental protocol, DART-ValueSelect, and its locked decision-value tests.",
        ),
        changes,
        "organization_map",
    )

    probabilistic = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("Probabilistic accuracy requires")))
    set_text(
        probabilistic,
        "Probabilistic credibility requires more than point MAE: CRPS, pinball loss, tail-weighted CRPS, interval coverage, calibration, spike probabilities, and joint-path diagnostics are retained. These metrics diagnose whether a DART distribution is trustworthy but do not select the economic winner. The 2026 DART study used as a comparator did not select models from annual mean DART; its probability and spike thresholds were chosen by validation P&L and its held-out evaluation reported cumulative and annual P&L, volatility, and Sharpe ratio. Its annual mean DART summaries primarily explained INC/DEC tendencies. The relevant limitation for this study is instead the user model: a virtual INC/DEC trader does not face the data-center deadlines, SLA, battery state, migration cost, bandwidth, and electricity-demand constraints modeled here.",
        changes,
        "related_work_corrected_2026_comparator",
    )
    gap = paragraph_after_heading(doc, "2.4 Research gap and positioning")
    set_text(
        gap,
        "There is no user-independent globally best DART model. The same forecast can have different value for a virtual trader, storage operator, and constrained data-center load. This study therefore separates probabilistic credibility, gate-time actionability, and realized counterfactual decision value; only the third layer determines the final static selection, subject to the first two layers and operational feasibility gates.",
        changes,
        "research_gap_decision_focus",
    )

    # Canonical sign and model interface.
    target_interface = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("The NLR calibration task estimates")))
    set_text(
        target_interface,
        "The NLR calibration task estimates compute-to-IT-power, facility-power, PUE, cooling, and weather/solar-response parameters on a chronological subset of referenced source records and freezes the selected parameter set or ensemble after validation. The NYISO forecasting task keeps official historical day-ahead LBMP and real-time LBMP records as independently issued targets with separate feature-availability masks, scalers, checkpoints, calibration residuals, and leaderboards. The paper's canonical derived spread is dart_da_minus_rt = da_lbmp - rt_lbmp. The inverse diagnostic rt_spread_rt_minus_da = rt_lbmp - da_lbmp is retained, and every processed partition must satisfy np.allclose(dart_da_minus_rt, -rt_spread_rt_minus_da, equal_nan=True). INC/DEC labels and the interpretation of increasing or decreasing day-ahead purchases follow the canonical DA-minus-RT sign throughout. Negative prices remain in the primary data, and future realized weather is excluded from deployable forecasts.",
        changes,
        "canonical_dart_sign",
    )

    # New methodological subsection after forecasting baselines and ablations.
    base_heading = paragraph_with_text(doc, "5.8 Forecasting baselines and ablations")
    paragraphs = doc.paragraphs
    base_index = next(
        index for index, paragraph in enumerate(paragraphs) if paragraph._p is base_heading._p
    )
    anchor = paragraphs[base_index + 1]
    blocks = [
        (MARKER, "Heading 2"),
        (
            "Layer 1, probabilistic credibility, reports CRPS, pinball loss, tail-weighted CRPS using thresholds learned only from the training period, interval and exceedance calibration, and joint scenario quality. These scores are diagnostics and admission checks; minimum MAE, RMSE, CRPS, or spike error never by itself names the preferred DART model.",
            "Normal",
        ),
        (
            "Layer 2, actionability, uses gate-time feasible flexible energy F(d,z,t), computed only from the task queue, deadlines, site capacity, migration bandwidth, battery state and limits, and other information available before the decision. AWCRPS(m) = sum[F(d,z,t) CRPS(m,d,z,t)] / sum[F(d,z,t)]. Future realized price, workload, or PV may not enter F. Economic-action Brier/F1 evaluate the user-specific increase/no-action/decrease decision after degradation, migration, transaction, settlement-asymmetry, minimum-dispatch, and SLA costs define time-varying no-action bands. Feasible Opportunity Ranking Loss evaluates only physically feasible space-time moves and weights ordering errors by transferable energy and realized net opportunity magnitude.",
            "Normal",
        ),
        (
            "Layer 3, decision value, sends each candidate's calibrated DART distribution through the same optimizer, constraints, user parameters, initial SOC, task queue, scenario counts, reduction method, random seeds, MIP gap, time limit, and rolling-update rule. First-stage day-ahead decisions are frozen before replay. Realized counterfactual DA/RT prices, workload, and PV are used only after the decision to recompute DA settlement, RT deviation settlement, battery degradation, migration, SLA, transaction, and other operating costs. Outputs are operational cost reduction, incremental objective value, prescriptive regret, downside risk, and Profit Capture Ratio; they are not observed commercial profit.",
            "Normal",
        ),
        (
            "DART-ValueSelect is fit only from rolling sample-out-of-operation results for 2022-2024 and is frozen before the once-only locked 2025 evaluation. A candidate is admitted only if it meets the engineered SLA floor, CVaR budget, maximum probability of underperforming the baseline, gate-time forecast-plus-solve latency, and zero-infeasibility requirement. Paired 7-day moving-block bootstrap is primary, with 14- and 28-day robustness blocks. A Model Confidence Set excludes statistically inferior candidates; among retained feasible candidates, the maximum 95% lower confidence bound of incremental operational value is selected. Near ties are resolved by lower regret, better tail calibration, lower latency, then simplicity. If candidates remain statistically indistinguishable, the paper reports a superior set and does not force a winner.",
            "Normal",
        ),
        (
            "The optional Contextual DART-ValueSelect extension uses only gate-time calendar, trailing 24/168-hour DART volatility, model disagreement, spike probabilities, interval widths, battery SOC, flexible/migratable workload, spare compute, and load/PV uncertainty. LightGBM Ranker, quantile regression, or regularized multinomial logit estimates risk-adjusted model value. A daily post-hoc best-model selector is an unattainable upper bound only and is never presented as deployable.",
            "Normal",
        ),
    ]
    for text, style in blocks:
        anchor = insert_paragraph_after(anchor, text, style)
    changes.append("new_section_5_9")

    scenario_reduction = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("The paper profile draws 100")))
    set_text(
        scenario_reduction,
        scenario_reduction.text.rstrip()
        + " DART candidate comparisons reuse common full-scenario draws and a common reduction map; candidate-specific distributions change the forecast input, not the experimental budget or solver treatment.",
        changes,
        "common_scenario_budget",
    )
    objective = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("The counterfactual objective minimizes")))
    set_text(
        objective,
        "The counterfactual objective minimizes simulated day-ahead procurement cost, expected real-time operating and imbalance cost, and a CVaR term. It also includes battery degradation, migration energy and bandwidth charges, PV curtailment, engineered SLA slack, transaction or other declared operating costs, and optional grid-facing ramp/concentration penalties. Realized replay uses actual historical DA and RT records plus the held-out realized workload and modeled PV path only after decisions are frozen. Because task revenue is not observed and is common across feasible candidates, the paper ranks equivalent operational-cost reduction and incremental objective improvement rather than inventing commercial revenue or realized profit:",
        changes,
        "realized_operational_value_boundary",
    )

    # Experimental protocol text.
    locked = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("The NLR physical twin and workload model use")))
    set_text(
        locked,
        locked.text.rstrip()
        + " DART-ValueSelect is estimated from strict rolling out-of-sample operating replays in 2022-2024 (expanding history, with deep models retrained no more frequently than the declared monthly schedule). Before entering 2025, candidate models, calibration, action thresholds, risk budgets, selector rules, user parameters, scenario counts, and inference rules are frozen. The locked 2025 period is evaluated once and cannot feed back into selection.",
        changes,
        "rolling_selection_and_locked_test",
    )
    set_text(
        paragraph_after_heading(doc, "8.3 Forecast metrics"),
        "The three reporting layers are: (1) probabilistic credibility--CRPS, pinball, tail-weighted CRPS, 50/80/90/95% coverage and width, PIT or exceedance calibration, spike PR-AUC, energy and variogram scores, temporal autocorrelation, cross-zone rank dependence, and spike duration/co-occurrence; (2) actionability--AWCRPS, economic-action Brier and weighted precision/recall/F1, false-action and missed-opportunity loss, and Feasible Opportunity Ranking Loss; and (3) decision value--paired realized operational cost, incremental objective value, regret, PCR, risk, service, and latency. Layers 1 and 2 are diagnostics and admission gates; layer 3 is the primary ranking evidence.",
        changes,
        "section_8_3_three_layers",
    )
    set_text(
        paragraph_after_heading(doc, "8.5 Optimization baselines"),
        "In addition to B0-B8, three DART experiments isolate selection value. A1 replaces only the candidate DART distribution while holding DA, load, PV, workload forecasts, dependence/scenario generator, optimizer, seeds, and all constraints fixed. A2 evaluates each candidate's complete forecast, calibration, and scenario chain under the same operational problem. A3 preregisters metric-contradiction cases in which the lowest MAE or CRPS does not produce the highest decision value, and reports their frequency and mechanism without selecting them after test inspection. Selection-method comparators are MAE, CRPS, spike F1, validation mean operational value, DART-ValueSelect, and the optional contextual selector.",
        changes,
        "section_8_5_a1_a3",
    )
    set_text(
        paragraph_after_heading(doc, "8.6 Decision and service metrics"),
        "For model m and day d, the frozen decision is replayed under the same realized historical DA/RT prices and held-out workload/PV realization to compute total operating cost and its value-equivalent Pi_real(m,d). Incremental value is Pi_real(m,d) - Pi_real(B,d); regret is Pi_PI(d) - Pi_real(m,d), where perfect information is an unattainable upper bound. PCR = sum_d[Pi_real(m,d) - Pi_real(B,d)] / sum_d[Pi_PI(d) - Pi_real(B,d)]: one captures the full perfect-information opportunity, zero equals the baseline, and a negative value underperforms it. Report cost reduction, value per MWh of flexibility and per GPU-hour, positive-value days, loss probability, CVaR95 of negative incremental value, worst-5% loss, maximum drawdown, maximum daily loss, consecutive-loss duration, SLA/deadline outcomes, and forecast-plus-solve latency. Sharpe ratio is supplementary only.",
        changes,
        "section_8_6_pcr_and_risk",
    )
    set_text(
        paragraph_after_heading(doc, "8.7 Ablations and sensitivity matrix"),
        "DART-ValueSelect uses the exact same physical and market problem for every candidate. Feasibility gates, the primary 7-day block bootstrap, 14/28-day blocks, MCS confidence level, risk budget, underperformance-probability budget, latency limit, and tie rules are frozen in configuration. Contextual selection is optional and cannot replace the static confirmatory result. Regime tables cover ordinary, positive-tail, negative-tail, high-workload, low-PV, and high-congestion days using thresholds fixed before the locked test.",
        changes,
        "section_8_7_selector_controls",
    )
    bootstrap = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("Report block-bootstrap")))
    set_text(
        bootstrap,
        "Report paired moving-block-bootstrap confidence intervals with a 7-day primary block and 14- and 28-day robustness blocks; all candidates use the same resampling indices so weekly dependence and value autocorrelation are preserved.",
        changes,
        "bootstrap_blocks",
    )
    dm = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("Use Diebold-Mariano-style")))
    set_text(
        dm,
        "Use Diebold-Mariano-style tests for forecast losses with horizon-aware variance, paired decision-value inference for operational outcomes, and Holm correction for preregistered contrasts. Apply a Model Confidence Set to daily realized operating cost or regret before selecting by the incremental-value lower confidence bound.",
        changes,
        "mcs_and_paired_inference",
    )
    no_test_select = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("Do not select a headline model")))
    set_text(
        no_test_select,
        "Do not select or tune a headline model from the locked 2025 cost, DART realization, perfect-information result, or any test-derived threshold. The selector is frozen from 2022-2024 rolling out-of-sample evidence and is then applied exactly once; a statistically indistinguishable superior set is an admissible result.",
        changes,
        "locked_selection_rule",
    )

    # Reporting text and mechanisms.
    set_text(
        paragraph_after_heading(doc, "9.1 NLR twin validation and NYISO DA/RT forecast tables"),
        "DA and RT tables remain separate, while the DART tables use only the canonical DA-minus-RT sign. The four added DART templates distinguish probability/actionability metrics, realized decision value, selection strategies, and market regimes. Every value remains TO BE GENERATED until linked to a complete hash-verified run.",
        changes,
        "reporting_dart_tables",
    )
    set_text(
        paragraph_after_heading(doc, "9.4 Counterfactual optimization main table"),
        "Report each candidate and selection strategy under the identical optimizer: operational cost reduction, incremental objective value, PCR, prescriptive regret, loss probability, CVaR95, maximum drawdown, SLA, latency, infeasibility, timeout, and MIP gap. Perfect information is labeled UNATTAINABLE UPPER BOUND and never compared as a deployable policy.",
        changes,
        "reporting_decision_value",
    )
    set_text(
        paragraph_after_heading(doc, "9.5 Hypothesis decision rules"),
        "H4 is decomposed into H4a-H4d so disagreement between statistical rank and constrained user value, the static ValueSelect benefit, explanatory value of tail/actionability diagnostics, and optional contextual selection are tested separately. Unsupported and indistinguishable outcomes remain reportable results.",
        changes,
        "reporting_h4_subhypotheses",
    )
    fig_anchor = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("Runtime and MIP-gap scaling")))
    insert_paragraph_after(
        fig_anchor,
        "DART credibility-actionability-value rank comparison; selection LCB and Model Confidence Set; metric-contradiction cases; and value by ordinary, tail, high-workload, low-PV, and congestion regimes.",
        "List Number",
    )
    changes.append("required_dart_figures")
    mechanism = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("Probabilistic forecasting may create")))
    set_text(
        mechanism,
        "Probabilistic forecasting may create counterfactual decision value even when median error changes little, and a lower unconditional error need not imply a better constrained action. Correctly representing NYISO DART tails, NLR workload bursts, and physical-parameter uncertainty can reduce simulated operating cost, regret, CVaR, or engineered service slack. The optimizer is therefore an economic stress test of forecast usefulness, not evidence that NLR actually participated in NYISO, and only realized replay under frozen decisions can support the counterfactual ranking.",
        changes,
        "mechanism_decision_focus",
    )
    conclusion = paragraph_with_text(doc, next(p.text for p in doc.paragraphs if p.text.startswith("DGPT-E uses dynamic graph")))
    set_text(
        conclusion,
        conclusion.text.rstrip()
        + " Canonical DA-minus-RT DART distributions are evaluated by DART-ValueSelect: probability and actionability metrics screen credibility, but the final selection uses risk-adjusted, rolling out-of-sample counterfactual user value from identical optimization and replay conditions.",
        changes,
        "conclusion_value_select",
    )

    # Hypothesis tables.
    hypotheses = find_table(doc, "ID", "Preregistered hypothesis")
    h4_row = next(row for row in hypotheses.rows if row.cells[0].text.strip() == "H4")
    h4_row.cells[0].text = "H4a"
    h4_row.cells[1].text = "Rankings from traditional MAE, CRPS, and spike-classification metrics differ materially from rankings based on realized counterfactual data-center decision value under identical constraints."
    insert_rows_before_label(
        hypotheses,
        "H5",
        [
            ["H4b", "The DART-ValueSelect model or superior set yields greater locked-test operational cost reduction and lower prescriptive regret than selection by MAE, CRPS, spike F1, or simple validation mean value, without breaching service or risk gates."],
            ["H4c", "Tail calibration, economic-action identification, and feasible opportunity ranking explain counterfactual decision value better than unconditional point error."],
            ["H4d", "The optional contextual selector improves risk-adjusted counterfactual value over the frozen static selector without increasing SLA violations or tail loss; if it is not run, H4d is explicitly NOT EVALUATED rather than inferred."],
        ],
    )
    changes.append("hypotheses_h4a_h4d")

    rules = find_table(doc, "Hypothesis", "Decision rule")
    h4_rule = next(row for row in rules.rows if row.cells[0].text.strip() == "H4")
    h4_rule.cells[0].text = "H4a"
    h4_rule.cells[1].text = "Supported only by preregistered paired rank-disagreement evidence between statistical metrics and realized counterfactual decision value; anecdotal metric contradictions are insufficient."
    insert_rows_before_label(
        rules,
        "H5",
        [
            ["H4b", "Supported only if the feasible MCS-retained DART-ValueSelect model/superior set improves locked-test operational cost or value and prescriptive regret versus every preregistered selection comparator without violating SLA, risk, latency, or feasibility gates."],
            ["H4c", "Supported only if preregistered tail-calibration, economic-action, and feasible-ranking diagnostics have stronger out-of-sample explanatory association with paired decision value than unconditional point error, with uncertainty and multiplicity control."],
            ["H4d", "Supported only if the optional contextual selector improves paired risk-adjusted locked-test value over the frozen static selector while noninferior on SLA and tail loss; otherwise NOT SUPPORTED or NOT EVALUATED."],
        ],
    )
    changes.append("hypothesis_rules_h4a_h4d")

    # Existing methodological tables.
    issues = find_table(doc, "Issue", "Revised treatment")
    append_row(issues, ["DART sign and model selection", "Use DART = DA - RT as the sole paper sign; retain RT - DA only as an explicitly named inverse diagnostic and test exact negation. Select models from rolling constrained decision value, never annual mean DART or minimum average error alone."])
    changes.append("issues_table_dart")

    doc_map = find_table(doc, "Section", "Purpose")
    for row in doc_map.rows:
        if row.cells[0].text.strip() == "5. Calibration and forecasting":
            row.cells[1].text = "NLR physical twin, independent NYISO DA/RT DGPT-E, NLR CfC-W, canonical DART derivation, calibration, and three-layer DART evaluation"
        elif row.cells[0].text.strip() == "8. Experimental protocol":
            row.cells[1].text = "Independent splits, A1-A3, DART-ValueSelect, baselines, ablations, transfer support, rolling tests, and inference"
        elif row.cells[0].text.strip() == "9. Reporting":
            row.cells[1].text = "Evidence-layer tables, DART probability/actionability/value/selection/regime tables, figures, hypothesis rules, and no-results templates"
    changes.append("document_map_dart")

    gaps = find_table(doc, "Gap", "Why it matters")
    append_row(gaps, ["DART model selection", "Minimum average forecast error can miss economically asymmetric tails, feasible rankings, and intertemporal user constraints.", "Compare every candidate in identical stochastic optimization and realized replay; filter by feasibility/risk and select from the MCS by incremental-value LCB."])
    changes.append("gap_table_dart")

    branches = find_table(doc, "Branch", "Targets")
    append_row(branches, ["DART decision layer", "Canonical DART = DA LBMP - RT LBMP, action probabilities, and feasible opportunity rankings derived from aligned candidate distributions", "Rolling 2022-2024 out-of-sample selection; once-only locked 2025 test", "Credibility/actionability metrics, common-seed scenarios, frozen decisions, realized replay, and ValueSelect evidence"])
    for row in branches.rows:
        if row.cells[0].text.strip() == "DGPT-E NYISO-RT":
            row.cells[1].text = "Zonal real-time LBMP; canonical DART (DA - RT) derived jointly; inverse RT-minus-DA spread retained only as a named diagnostic"
    changes.append("forecast_branches_dart")

    models = find_table(doc, "Group", "Models or removals")
    append_row(models, ["DART selectors", "MAE, CRPS, spike F1, validation mean operational value, DART-ValueSelect; optional Contextual DART-ValueSelect with LightGBM Ranker, quantile regression, or regularized multinomial logit"])
    changes.append("models_table_selectors")

    diagnostics = find_table(doc, "Diagnostic family", "Required outputs")
    append_row(diagnostics, ["DART actionability", "AWCRPS, economic-action Brier/weighted precision-recall-F1, false-action and missed-opportunity loss, Feasible Opportunity Ranking Loss"])
    changes.append("scenario_diagnostics_actionability")

    forecast_metrics = find_table(doc, "Family", "Metrics")
    append_row(forecast_metrics, ["DART credibility and actionability", "Canonical-sign DART CRPS, tail-CRPS, AWCRPS, 50/80/90/95% coverage and width, PIT/exceedance calibration, action Brier/F1, spike PR-AUC, FORL, energy/variogram scores"])
    changes.append("forecast_metrics_dart")

    decision_metrics = find_table(doc, "Family", "Required metrics")
    append_row(decision_metrics, ["DART selection value", "Paired operational cost reduction, incremental objective value, correct PCR, prescriptive regret, loss probability, CVaR95, worst-5% loss, maximum drawdown, SLA, forecast-plus-solve latency, infeasibility and MIP gap"])
    changes.append("decision_metrics_dart")

    sensitivity = find_table(doc, "Dimension", "Levels")
    append_row(sensitivity, ["DART selection", "A1 replacement/A2 full chain/A3 contradiction; 7-day primary and 14/28-day bootstrap; MCS; static selector; optional contextual selector; ordinary/tail/workload/PV/congestion regimes"])
    changes.append("sensitivity_dart")

    # Four required result templates inserted after the optimization template.
    optimization_template = find_table(doc, "Policy", "Mean cost")
    table_style = optimization_template.style.name if optimization_template.style else "Table Grid"
    anchor_element = optimization_template._tbl
    placeholder = "TO BE GENERATED"
    model_rows = [
        [model] + [placeholder] * 7
        for model in ["Seasonal naive", "LightGBM", "PatchTST", "iTransformer", "DGPT-E"]
    ]
    anchor_element = add_template_table_after(
        doc,
        anchor_element,
        "DART Table 1. Probabilistic credibility and actionability (canonical DA-minus-RT sign)",
        ["Model", "CRPS", "Tail-CRPS", "AWCRPS", "Action Brier", "Spike PR-AUC", "Coverage", "Energy score"],
        model_rows,
        table_style,
    )
    decision_rows = [
        [model] + [placeholder] * 8
        for model in ["Baseline", "LightGBM", "PatchTST", "iTransformer", "DGPT-E", "DART-ValueSelect", "Perfect information (upper bound only)"]
    ]
    anchor_element = add_template_table_after(
        doc,
        anchor_element,
        "DART Table 2. Counterfactual user decision value under identical optimization and realized replay",
        ["Model", "Incremental value", "Cost reduction", "PCR", "Regret", "Loss probability", "CVaR95", "Max drawdown", "SLA"],
        decision_rows,
        table_style,
    )
    selection_rows = [
        [method] + [placeholder] * 7
        for method in ["MAE selection", "CRPS selection", "Spike-F1 selection", "Validation mean-value selection", "DART-ValueSelect", "Contextual DART-ValueSelect (optional)"]
    ]
    anchor_element = add_template_table_after(
        doc,
        anchor_element,
        "DART Table 3. Comparison of model-selection strategies",
        ["Selection method", "Selected model/set", "Locked-test value", "PCR", "Regret", "CVaR95", "SLA", "Inference + solve time"],
        selection_rows,
        table_style,
    )
    regime_rows = [
        [model] + [placeholder] * 6
        for model in ["Static DART-ValueSelect", "Contextual DART-ValueSelect (optional)", "Best single candidate", "Post-hoc daily oracle (upper bound only)"]
    ]
    add_template_table_after(
        doc,
        anchor_element,
        "DART Table 4. Counterfactual incremental value by preregistered market regime",
        ["Model/selector", "Ordinary day", "Positive-tail day", "Negative-tail day", "High-workload day", "Low-PV day", "High-congestion day"],
        regime_rows,
        table_style,
    )
    changes.append("four_dart_result_tables")

    # Symbols, processed data, and artifact mapping.
    symbols = find_table(doc, "Symbol", "Meaning")
    for row in [
        ["DART(z,t)", "Canonical day-ahead minus real-time LBMP: lambda_DA(z,t) - lambda_RT(z,t)"],
        ["F(d,z,t)", "Gate-time feasible flexible energy; never computed with future realized prices"],
        ["Pi_real(m,d)", "Counterfactual realized operational value-equivalent after frozen-decision replay"],
        ["R(m,d)", "Prescriptive regret: perfect-information upper-bound value minus model value"],
        ["PCR(m)", "Captured share of perfect-information opportunity relative to the same baseline"],
        ["LCB(m)", "95% lower confidence bound of paired out-of-sample incremental operational value"],
    ]:
        append_row(symbols, row)
    changes.append("symbols_dart")

    data_fields = find_table(doc, "File", "Required fields")
    market_row = next(row for row in data_fields.rows if row.cells[0].text.strip() == "nyiso_market_15m.parquet")
    market_row.cells[1].text = "timestamp_utc/local, zone, da_lbmp, rt_lbmp, dart_da_minus_rt, rt_spread_rt_minus_da, sign_identity_passed, load, forecast, source_resolution, availability timestamp, missing flags, source_relationship, evidence_layer"
    append_row(data_fields, ["dart_value_select_daily.parquet", "date, model, selector, split, frozen-decision hash, realized-replay hash, probability/actionability metrics, cost components, incremental value, regret, PCR denominator, risk, SLA, latency, feasibility, MIP gap, bootstrap/MCS IDs, evidence layer"])
    changes.append("processed_fields_dart")

    artifacts = find_table(doc, "Artifact", "Paper destination")
    for row in [
        ["artifacts/tables/dart_probability_quality.csv", "Section 5.9/9.1; H4a-H4c"],
        ["artifacts/tables/dart_decision_value.csv", "Sections 8.6/9.4; H4b-H4d"],
        ["artifacts/tables/dart_selection_methods.csv", "Sections 5.9/9.4; H4a-H4b"],
        ["artifacts/tables/dart_regime_value.csv", "Sections 5.9/9.4; H4d and mechanism analysis"],
        ["artifacts/tables/dart_metric_contradictions.csv", "Experiment A3; H4a-H4c"],
        ["artifacts/manifests/dart_value_select_manifest.json", "Selection freeze, feasibility, common-optimizer, bootstrap, MCS, locked-test and claim-boundary evidence"],
    ]:
        append_row(artifacts, row)
    changes.append("artifact_mapping_dart")

    # Integrity checklist.
    checklist_anchor = paragraph_with_text(
        doc,
        next(p.text for p in doc.paragraphs if "final manuscript preserves" in p.text),
    )
    for text in [
        "☐ Every partition satisfies dart_da_minus_rt = -rt_spread_rt_minus_da with missing values aligned; all INC/DEC and DA-purchase direction labels use canonical DA-minus-RT semantics.",
        "☐ DART model selection uses only rolling 2022-2024 sample-out-of-operation evidence; all selector settings are frozen before the once-only locked 2025 test.",
        "☐ Every DART candidate receives identical user constraints, initial state, scenario budget/reduction, common random seeds, solver time/MIP gap, and rolling replay rules; A1 and A2 are not conflated.",
        "☐ PCR uses (model - baseline)/(perfect information - baseline); perfect information and post-hoc daily selection are labeled unattainable upper bounds.",
        "☐ No invented revenue or assumed commercial profit is reported; claims use counterfactual operational cost reduction, incremental objective value, regret reduction, and PCR.",
    ]:
        checklist_anchor = insert_paragraph_after(checklist_anchor, text, "Normal")
    changes.append("integrity_checklist_dart")

    doc.save(target)
    after_hash = sha256(target)
    after = package_snapshot(target)
    with zipfile.ZipFile(target) as archive:
        bad_member = archive.testzip()
    if bad_member is not None:
        raise RuntimeError(f"Corrupt DOCX member after save: {bad_member}")

    required_phrases = [
        MARKER,
        "dart_da_minus_rt = da_lbmp - rt_lbmp",
        "rt_spread_rt_minus_da = rt_lbmp - da_lbmp",
        "PCR = sum_d[Pi_real(m,d) - Pi_real(B,d)] / sum_d[Pi_PI(d) - Pi_real(B,d)]",
        "H4d",
        "DART Table 4. Counterfactual incremental value by preregistered market regime",
    ]
    check_doc = Document(target)
    full_text = "\n".join(p.text for p in check_doc.paragraphs)
    full_text += "\n" + "\n".join(
        "\t".join(cell.text for row in table.rows for cell in row.cells)
        for table in check_doc.tables
    )
    missing = [phrase for phrase in required_phrases if phrase not in full_text]
    if missing:
        raise RuntimeError(f"Required manuscript phrases missing: {missing}")
    if after["sections"] != before["sections"] or after["inline_shapes"] != before["inline_shapes"]:
        raise RuntimeError("Unexpected loss of sections or inline shapes")

    manifest = {
        "status": "APPLIED",
        "applied_at_utc": datetime.now(timezone.utc).isoformat(),
        "target": str(target),
        "backup": str(backup),
        "input_sha256": original_hash,
        "backup_sha256": sha256(backup),
        "output_sha256": after_hash,
        "structure_before": before,
        "structure_after": after,
        "changes": changes,
        "required_phrase_checks": {phrase: True for phrase in required_phrases},
        "zip_integrity": "PASS",
        "render_or_visual_review_performed": False,
        "render_skip_reason": "Explicit user and goal instruction: no manuscript rendering, layout, formatting, or manual visual review.",
    }
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
