from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

from docx import Document


OLD_TEXT = (
    "Use one global random seed list fixed before experiments. Under the 2026-08-11 "
    "protocol amendment, every confirmatory training, replay, optimization, sensitivity, "
    "and paired-inference panel uses exactly [20260809, 20260819, 20260829, 20260908]. "
    "Seed 20260918 is excluded before its pending CfC-W trial and formal DART training; "
    "any earlier artifact containing it is retained only as "
    "legacy_pre_amendment_excluded_from_confirmatory_four_seed_panel. Report the "
    "four-seed mean, standard deviation, and every failed seed."
)

PREVIOUS_THREE_TEXT = (
    "Use one global random seed list fixed before experiments. The 2026-08-11 three-seed "
    "protocol amendment supersedes the earlier four-seed amendment for every experiment "
    "that starts afterward. Forecast extensions, common-DA-gate DART training and A1/A2 "
    "replay, optimization, sensitivity, and paired-inference panels use exactly "
    "[20260809, 20260819, 20260829]. The completed four-seed paper-forecast-v1 recovery "
    "remains immutable auxiliary evidence; seed 20260908 is retained for provenance but "
    "excluded from every subsequent three-seed completion count, mean, uncertainty "
    "estimate, ranking, winner set, and paper claim. Seed 20260918 remains excluded and "
    "was not started for the pending CfC-W or formal DART stages. Label both excluded "
    "lineages legacy_pre_amendment_excluded_from_confirmatory_three_seed_panel. DART "
    "annual training contains 168 logical units (144 neural, 8 LightGBM, and 16 "
    "deterministic), forecast extensions contain 63 seed-runs, and the top-level DART "
    "replay ledger contains 84,738 cells. Report the three-seed mean, standard deviation, "
    "and every failed seed."
)

NEW_TEXT = PREVIOUS_THREE_TEXT + (
    " For every native readiness panel and fixed 96-step forecast bundle, operating_date "
    "is the canonical UTC day with exactly 96 15-minute intervals; America/New_York is "
    "retained as NYISO market-time and DST metadata rather than the grouping clock."
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    repository = Path(__file__).resolve().parents[1]
    workspace = repository.parent
    source = workspace / (
        "Data_Center_VPP_English_Manuscript_and_Experimental_Protocol_"
        "NLR_NYISO_Revised.docx"
    )
    backup = repository / "artifacts" / "manifests" / "specification_backups" / (
        "Data_Center_VPP_English_Manuscript_and_Experimental_Protocol_"
        "NLR_NYISO_Revised.pre_three_seed_amendment_20260811.docx"
    )
    manifest = repository / "artifacts" / "manifests" / (
        "three_seed_protocol_amendment.json"
    )
    backup.parent.mkdir(parents=True, exist_ok=True)
    manifest.parent.mkdir(parents=True, exist_ok=True)

    document = Document(source)
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text in {OLD_TEXT, PREVIOUS_THREE_TEXT}
    ]
    already = [paragraph for paragraph in document.paragraphs if paragraph.text == NEW_TEXT]
    if len(already) == 1 and not matches:
        return
    if len(matches) != 1 or already:
        raise RuntimeError(
            f"expected one old paragraph and no new paragraph; got old={len(matches)}, new={len(already)}"
        )

    original_sha256 = sha256(source)
    if not backup.exists():
        shutil.copy2(source, backup)

    paragraph = matches[0]
    if not paragraph.runs:
        paragraph.add_run(NEW_TEXT)
    else:
        paragraph.runs[0].text = NEW_TEXT
        for run in paragraph.runs[1:]:
            run.text = ""

    temporary = source.with_suffix(".three-seed.tmp.docx")
    document.save(temporary)
    verified = Document(temporary)
    if sum(paragraph.text == NEW_TEXT for paragraph in verified.paragraphs) != 1:
        temporary.unlink(missing_ok=True)
        raise RuntimeError("three-seed protocol paragraph did not persist exactly once")
    temporary.replace(source)

    payload = {
        "schema_version": "dcvpp.three_seed_protocol_amendment/v1",
        "status": "COMPLETE",
        "render_or_layout_check_performed": False,
        "render_or_layout_check_reason": "user_explicitly_requested_no_format_checks",
        "source_path": str(source.relative_to(workspace)),
        "source_sha256_before": original_sha256,
        "source_sha256_after": sha256(source),
        "backup_path": str(backup.relative_to(repository)),
        "backup_sha256": sha256(backup),
        "confirmatory_seed_panel": [20260809, 20260819, 20260829],
        "excluded_legacy_seeds": [20260908, 20260918],
        "completed_v1_four_seed_evidence_preserved": True,
        "dart_annual_training_units": {
            "neural": 144,
            "lightgbm": 8,
            "deterministic": 16,
            "total": 168,
        },
        "forecast_extension_seed_runs": 63,
        "dart_top_level_replay_cells": {
            "selection": 59184,
            "locked_2025": 19710,
            "shared_baseline": 4383,
            "perfect_information": 1461,
            "total": 84738,
        },
    }
    temporary_manifest = manifest.with_suffix(manifest.suffix + ".tmp")
    temporary_manifest.write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    temporary_manifest.replace(manifest)


if __name__ == "__main__":
    main()
