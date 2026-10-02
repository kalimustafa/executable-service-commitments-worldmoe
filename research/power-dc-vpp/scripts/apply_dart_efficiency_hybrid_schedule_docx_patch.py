"""Apply the efficiency-first hybrid DART scheduling amendment to the manuscript.

The user explicitly waived render/layout QA. This script preserves the existing
document structure and performs text, hash, and reopen checks only.
"""

from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import shutil

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
REPO = Path(__file__).resolve().parents[1]
DOCX = ROOT / "Data_Center_VPP_English_Manuscript_and_Experimental_Protocol_NLR_NYISO_Revised.docx"
SPEC = ROOT / "CODEX_Executable_Experiment_Specification_English_NLR_NYISO_Revised.md"
ADAPTER = REPO / "docs" / "dart_training_adapter.md"
BACKUP = (
    REPO
    / "artifacts"
    / "manifests"
    / "specification_backups"
    / "Data_Center_VPP_English_Manuscript_and_Experimental_Protocol_NLR_NYISO_Revised.pre_efficiency_hybrid_schedule_20260812.docx"
)
MANIFEST = REPO / "artifacts" / "manifests" / "dart_efficiency_hybrid_schedule_document_amendment.json"

EXPECTED_SPEC_SHA256_BEFORE = "4a9f4b1f012754421b72ebf7f7dde24cdaffb970d302269d0471faadb966abf7"
EXPECTED_ADAPTER_SHA256_BEFORE = "fc48fa65b480ce166b0000afd8143b029cb2c613f8210ab5e6f8ca48f2155189"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def find_exact(doc: Document, text: str):
    matches = [paragraph for paragraph in doc.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise RuntimeError(f"expected one exact paragraph, found {len(matches)}: {text[:96]}")
    return matches[0]


def main() -> None:
    for path in (DOCX, SPEC, ADAPTER):
        if not path.exists():
            raise FileNotFoundError(path)

    spec_after = sha256_file(SPEC)
    adapter_after = sha256_file(ADAPTER)
    spec_text = SPEC.read_text(encoding="utf-8")
    adapter_text = ADAPTER.read_text(encoding="utf-8")
    required_spec_fragments = (
        "Efficiency-first hybrid cross-model execution amendment",
        "strict_cross_model_rotation_every_5_epochs: false",
        "global_assessment_epochs: [40, 50, 60, 70, 80, 90, 100]",
        "Individual-seed scheduling, early stopping, and per-seed epoch selection are forbidden",
        "Primary-before-ablation execution amendment",
    )
    required_adapter_fragments = (
        "efficiency-first hybrid barrier schedule",
        "one uninterrupted\n10-epoch increment per global round",
        "mid-block audit\ndoes not force a model switch or materialize a stop",
        "Structural ablations and ORACLE variants remain deferred",
    )
    missing_spec = [fragment for fragment in required_spec_fragments if fragment not in spec_text]
    missing_adapter = [fragment for fragment in required_adapter_fragments if fragment not in adapter_text]
    if missing_spec or missing_adapter:
        raise RuntimeError(
            f"Markdown amendment check failed: spec={missing_spec}, adapter={missing_adapter}"
        )

    before_sha = sha256_file(DOCX)
    doc = Document(DOCX)
    marker = "efficiency-first hybrid neural scheduler"
    if any(marker in paragraph.text.lower() for paragraph in doc.paragraphs):
        raise RuntimeError("efficiency-first hybrid scheduling amendment is already present")

    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if BACKUP.exists() and sha256_file(BACKUP) != before_sha:
        raise RuntimeError("existing pre-amendment backup differs from the current input")
    if not BACKUP.exists():
        shutil.copy2(DOCX, BACKUP)

    old = (
        "Use three-seed group-synchronized early stopping: no stop before epoch 40; assess every 5 epochs; stop all "
        "three seeds together only after three consecutive stable assessments of the selected epoch and both proper-score "
        "survivor sets with no material proper-score improvement; stop at epoch 100 at the latest. Individual-seed "
        "early stopping and per-seed epoch selection are forbidden."
    )
    new = (
        "Use the efficiency-first hybrid neural scheduler. Treat each candidate by annual origin by DA/RT target head "
        "and its exact three seeds as one synchronized group. Advance one group continuously through epoch 40 before "
        "moving to the next group, and perform the first common assessment only after every required group has reached "
        "epoch 40. Thereafter, advance every active group by one uninterrupted 10-epoch increment per global round and "
        "make stop decisions only after all active groups reach the shared boundaries at epochs 50, 60, 70, 80, 90, "
        "and 100. Preserve a 5-epoch checkpoint, health, and provisional-score ledger inside each block, but do not "
        "switch models or stop a group at the mid-block audit. Remove a group from later rounds only at a global "
        "boundary when the frozen three-consecutive-stable-assessment rule passes. All three seed replicas stop "
        "together; individual-seed scheduling, stopping, and checkpoint selection are forbidden. The schedule is "
        "fixed for efficiency rather than candidate performance, and epoch 100 is the hard maximum."
    )
    paragraph = find_exact(doc, old)
    original_style = paragraph.style.name
    paragraph.text = new
    paragraph.style = original_style
    doc.save(DOCX)

    after_sha = sha256_file(DOCX)
    reopened = Document(DOCX)
    all_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    required_docx_fragments = (
        "efficiency-first hybrid neural scheduler",
        "first common assessment only after every required group has reached epoch 40",
        "one uninterrupted 10-epoch increment per global round",
        "5-epoch checkpoint, health, and provisional-score ledger",
        "individual-seed scheduling, stopping, and checkpoint selection are forbidden",
        "primary-before-ablation",
    )
    missing_docx = [fragment for fragment in required_docx_fragments if fragment not in all_text]
    if missing_docx:
        raise RuntimeError(f"post-save structural text check failed: {missing_docx}")
    if old in all_text:
        raise RuntimeError("superseded strict five-epoch stopping paragraph remains in the document")

    manifest = {
        "schema_version": "dcvpp.dart_efficiency_hybrid_schedule_document_amendment/v1",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "document_path": str(DOCX),
        "document_sha256_before": before_sha,
        "document_sha256_after": after_sha,
        "backup_path": str(BACKUP),
        "backup_sha256": sha256_file(BACKUP),
        "authoritative_spec_path": str(SPEC),
        "authoritative_spec_sha256_before": EXPECTED_SPEC_SHA256_BEFORE,
        "authoritative_spec_sha256_after": spec_after,
        "training_adapter_path": str(ADAPTER),
        "training_adapter_sha256_before": EXPECTED_ADAPTER_SHA256_BEFORE,
        "training_adapter_sha256_after": adapter_after,
        "structural_checks_passed": True,
        "render_or_layout_qa_performed": False,
        "render_qa_omission_reason": "user_explicitly_requested_no_format_check",
        "training_or_experiment_started_by_this_amendment": False,
        "changes": [
            "supersede_strict_five_epoch_cross_group_rotation",
            "continuous_three_seed_group_warmup_to_epoch_40",
            "common_epoch_40_barrier",
            "ten_epoch_active_group_global_rounds",
            "five_epoch_ledger_without_mid_round_stop",
            "global_boundary_only_synchronized_stopping",
            "primary_nine_candidates_before_ablations_and_oracles",
        ],
    }
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({**manifest, "manifest_sha256": sha256_file(MANIFEST)}, indent=2))


if __name__ == "__main__":
    main()
