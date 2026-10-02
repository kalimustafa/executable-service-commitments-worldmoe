"""Apply the DART v12 anchor-only amendment without rendering the DOCX."""

from __future__ import annotations

from pathlib import Path
import shutil
import zipfile

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "Data_Center_VPP_English_Manuscript_and_Experimental_Protocol_NLR_NYISO_Revised.docx"
BACKUP = (
    Path(__file__).resolve().parents[1]
    / "artifacts/manifests/specification_backups"
    / "Data_Center_VPP_English_Manuscript_and_Experimental_Protocol_NLR_NYISO_Revised.pre_anchor_only_v12_20260812.docx"
)


REPLACEMENTS = {
    "Fit scalers, graph statistics, and q01/q05/q95/q99 tail thresholds on the core training history only. For every neural DART candidate, annual origin, and DA or RT head, align the complete epoch ledger across seeds 20260809, 20260819, and 20260829, and average each daily checkpoint metric across the three seeds before inference.":
        "Fit scalers, graph statistics, and q01/q05/q95/q99 tail thresholds on the core training history only. For every neural DART candidate, annual origin, and DA or RT head, align the exact common positive five-epoch anchor ledger across seeds 20260809, 20260819, and 20260829, and average each daily checkpoint metric across the three seeds before inference. Every completed v4 epoch also writes a full immutable audit/resume checkpoint, but non-five-epoch checkpoints never enter formal selection.",
    "Select the checkpoint through health and simultaneous 80%/90% calibration gates, followed by a one-sided simultaneous 95% global-CRPS confidence set and a proper-tail-twCRPS confidence set using 7-day circular moving blocks. Select the earliest surviving epoch; do not minimize one scalar score or use an arbitrary weighted composite.":
        "Select only from exact common positive five-epoch anchors persisted and hash-verified for all three seeds. Apply health and simultaneous 80%/90% calibration gates, followed by a one-sided simultaneous 95% global-CRPS confidence set and a proper-tail-twCRPS confidence set using 7-day circular moving blocks. Select the earliest surviving anchor directly; between-anchor replay, scalar-score minimization, and arbitrary weighted composites are forbidden.",
    "This checkpoint rule uses only chronologically pre-cutoff predictive evidence. Action metrics, MILP solutions, operating cost, profit, regret, CVaR, rolling evaluation-year outcomes, and locked-2025 targets are forbidden. It selects an epoch within one model family and does not change the outer minimum-MAE versus DART-ValueSelect comparison.":
        "This checkpoint rule uses only chronologically pre-cutoff predictive evidence. Action metrics, MILP solutions, operating cost, profit, regret, CVaR, rolling evaluation-year outcomes, and locked-2025 targets are forbidden. It selects an anchor within one model family and does not change architecture, loss, hyperparameters, data, splits, seeds, or the outer minimum-MAE versus DART-ValueSelect comparison.",
    "Use the efficiency-first hybrid neural scheduler. Treat each candidate by annual origin by DA/RT target head and its exact three seeds as one synchronized group. Advance one group continuously through epoch 40 before moving to the next group, and perform the first common assessment only after every required group has reached epoch 40. Thereafter, advance every active group by one uninterrupted 10-epoch increment per global round and make stop decisions only after all active groups reach the shared boundaries at epochs 50, 60, 70, 80, 90, and 100. Preserve a 5-epoch checkpoint, health, and provisional-score ledger inside each block, but do not switch models or stop a group at the mid-block audit. Remove a group from later rounds only at a global boundary when the frozen three-consecutive-stable-assessment rule passes. All three seed replicas stop together; individual-seed scheduling, stopping, and checkpoint selection are forbidden. The schedule is fixed for efficiency rather than candidate performance, and epoch 100 is the hard maximum.":
        "Use the efficiency-first hybrid neural scheduler. Treat each candidate by annual origin by DA/RT target head and its exact three seeds as one synchronized group. Advance one group continuously through epoch 40 before moving to the next group, and perform the first common assessment only after every required group has reached epoch 40. Thereafter, advance every active group by one uninterrupted 10-epoch increment per global round and make stop decisions only after all active groups reach the shared boundaries at epochs 50, 60, 70, 80, 90, and 100. Under v4, every completed epoch writes an append-only hash-pinned full checkpoint containing model, optimizer, scaler, Python/NumPy/Torch/CUDA RNG, DataLoader generator/sampler, health, and validation-ledger state. Formal provisional and final selection still use only the exact common positive five-epoch anchors; intervening checkpoints are audit/resume only. Remove a group from later rounds only at a global boundary when the frozen three-consecutive-stable-assessment rule passes. All three seed replicas stop together; individual-seed scheduling, stopping, checkpoint selection, and between-anchor replay are forbidden. The schedule is fixed for efficiency rather than candidate performance, and epoch 100 is the hard maximum.",
    "The within-model checkpoint confidence sets are separate from the outer A2 Model Confidence Set. The former use only pre-cutoff predictive evidence and choose the earliest statistically credible epoch; the latter uses 2022-2024 causal market replay to form the DART-ValueSelect winner set.":
        "The within-model checkpoint confidence sets are separate from the outer A2 Model Confidence Set. The former use only pre-cutoff predictive evidence on the exact common positive five-epoch anchor grid and choose the earliest statistically credible persisted anchor; the latter uses 2022-2024 causal market replay to form the DART-ValueSelect winner set.",
}


def main() -> None:
    doc = Document(DOCX)
    if any("paper-dart-v1-readiness-v8" in paragraph.text for paragraph in doc.paragraphs):
        raise RuntimeError("v12 DOCX amendment already present")
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(DOCX, BACKUP)
    paragraph_count = len(doc.paragraphs)
    table_count = len(doc.tables)
    for old, new in REPLACEMENTS.items():
        matches = [paragraph for paragraph in doc.paragraphs if paragraph.text == old]
        if len(matches) != 1:
            raise RuntimeError(f"expected one paragraph, found {len(matches)}: {old[:72]}")
        matches[0].text = new
    anchor = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.startswith("Use the efficiency-first hybrid neural scheduler.")
    )
    amendment = doc.add_paragraph(
        "The paper-dart-forecast-v3 run failed closed before any OOS unit was materialized because CUDA replay of an arbitrary selected epoch changed its validation prediction or health record. Preserve v3 and the v7 readiness chain byte-for-byte. Continue under paper-dart-v1-readiness-v8, paper-dart-forecast-v4, paper-dart-replay-v4, and paper-dart-value-select-v4. Import v3 positive five-epoch anchors only through a read-only snapshot that re-hashes every source file; do not copy, hard-link, rewrite, or import mutable latest.pt, epoch-zero, selected-checkpoint, group, coordinator, or scheduler aliases. This reproducibility amendment consumed no 2022-2024 OOS outcome, locked-2025 result, replay, MILP, action, or value evidence.",
        style=anchor.style,
    )
    anchor._p.addnext(amendment._p)
    doc.save(DOCX)
    reopened = Document(DOCX)
    if len(reopened.paragraphs) != paragraph_count + 1 or len(reopened.tables) != table_count:
        raise RuntimeError("DOCX structure changed unexpectedly")
    if sum("paper-dart-v1-readiness-v8" in p.text for p in reopened.paragraphs) != 1:
        raise RuntimeError("v12 DOCX marker count is not one")
    with zipfile.ZipFile(DOCX) as package:
        if package.testzip() is not None:
            raise RuntimeError("DOCX ZIP integrity failed")


if __name__ == "__main__":
    main()
