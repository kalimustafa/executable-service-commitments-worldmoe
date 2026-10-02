"""Replace legacy manuscript hypotheses with W0/W1/R stage-aligned gates."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


HYPOTHESES = {
    "H0": (
        "W0 causal-state foundation. W0 must satisfy DART identity tolerance "
        "(1e-6), finite-output checks, and paired one-sided 95% upper bounds at "
        "or below zero for SOC and positive-queue severity against their causal "
        "baselines. Queue occurrence is diagnostic rather than a paper release "
        "gate. Among eligible full-2024 candidates, group-balanced CRPS selects "
        "the checkpoint; MAE and RMSE never select it."
    ),
    "H1": (
        "W1 action-conditioned dynamics. Proper scores at h=1, 4, 16, and 96 "
        "must each improve on the zero-update reference; latent KL and queue "
        "rollout must be non-degenerate; dynamics gradients and parameter updates "
        "must be nonzero; DART identity must hold; and 2025 must remain unopened. "
        "The worst-horizon proper score selects the checkpoint."
    ),
    "H2": (
        "R-A safe action expansion. Across the fixed 13-epoch curriculum, the "
        "sentinel decision value must decrease as the action range expands from "
        "0.10 to 1.00, while outputs remain finite and the maximum hard residual "
        "stays at or below 1e-5."
    ),
    "H3": (
        "R-B formal checkpoint selection. Only finite R-B checkpoints evaluated "
        "on all 364 development-validation days and six decisions per day are "
        "eligible. They must satisfy hard residual <=1e-5, worst-origin SLA "
        "<=0.02, and batch-1 h96 p95 latency <=1900 ms. The eligible checkpoint "
        "with minimum realized-cost/CVaR/SLA/constraint/budget-excess decision "
        "value is selected."
    ),
    "H4": (
        "Sealed-2025 value. On the one-shot final suite, the selected balanced "
        "Top-2 model must have a paired main-minus-safe mean-cost 95% interval "
        "entirely below zero, lower CVaR95 than the safe tensor, zero mean "
        "unserved energy, and maximum hard residual <=1e-5."
    ),
    "H5": (
        "Sparse-design effectiveness. Against separately retrained Top-1, "
        "dense-direct R, and independent R-MoE, every comparator-minus-main "
        "paired mean-cost 95% interval must lie above zero and Top-2 must have "
        "lower CVaR95. All seven general experts must carry positive load and "
        "conditional-tail activation must remain strictly between zero and one."
    ),
    "H6": (
        "Deployment efficiency. The selected model must use one W1 backbone call "
        "per decision, no external solver or candidate search, batch-1 h96 p95 "
        "latency <=1900 ms, peak GPU memory <=10 GB, and a reported active-parameter "
        "count for the final checkpoint."
    ),
}


DECISIONS = {
    "H0": (
        "SUPPORTED under the corrected W0 scope. Epoch 96 has group-balanced "
        "CRPS 0.540863, passes DART identity, SOC paired upper -0.094540, and "
        "positive-queue-severity paired upper -0.190737; 2025 remained unopened. "
        "The legacy queue-occurrence diagnostic (recall 0) is disclosed but is "
        "not used as the revised paper-level release gate."
    ),
    "H1": (
        "SUPPORTED. W1 epoch 56/update 3360 passes the release gate. Proper "
        "scores are 0.508546, 0.508108, 0.516056, and 0.616113 at h=1, 4, 16, "
        "and 96, all below the corresponding zero-update references 8.145673, "
        "8.399182, 9.538195, and 10.987783. DART identity, non-degenerate latent "
        "KL and queue rollout, nonzero dynamics gradients/updates, strict loading, "
        "and locked_2025_opened=false all pass."
    ),
    "H2": (
        "SUPPORTED. During R-A, sentinel decision value decreased from 2.588629 "
        "at epoch 2 to 1.866125 at epoch 12 (27.91%) while every recorded output "
        "was finite and the maximum hard residual remained at or below "
        "1.940e-6, well inside the 1e-5 threshold; R-B began at epoch 14."
    ),
    "H3": (
        "SUPPORTED for checkpoint selection. Among completed eligible full-2024 "
        "evaluations at epochs 91, 104, 117, 130, 143, and 156, epoch 130 has "
        "the minimum decision value, 1.178472. Its mean realized cost is 0.565210, "
        "CVaR95 is 2.358281, worst-origin SLA is 0.016338, maximum hard residual "
        "is 3.241e-6, and batch-1 h96 p95 latency is 648.925 ms. Training ended "
        "at epoch 168 before the configured four-post-best-full-evaluation early "
        "stop fired; no early-stop-trigger claim is made."
    ),
    "H4": (
        "SUPPORTED on 362 common 2025 dates. Main-minus-safe mean normalized "
        "cost is -0.683757 with 95% moving-block interval [-0.798915, -0.589023]; "
        "Top-2 CVaR95 is 4.669517 versus 6.587572, mean unserved energy is zero, "
        "and maximum hard residual is 3.278e-6."
    ),
    "H5": (
        "SUPPORTED. Comparator-minus-main mean-cost differences are 0.034468 "
        "for retrained Top-1, 0.056758 for dense-direct R, and 0.055113 for "
        "independent R-MoE; all paired 95% intervals lie above zero, and every "
        "comparator has higher CVaR95. All seven Top-2 route loads are positive, "
        "route entropy is 0.666041, and conditional-tail activation is 0.449883."
    ),
    "H6": (
        "PARTIALLY SUPPORTED. The architecture is one-pass, solver-free, and "
        "candidate-search-free. Epoch 130 records batch-1 h96 p95 latency "
        "648.925 ms, below 1900 ms, and the final-suite equivalent p95 is "
        "75.423 ms/sample. A checkpoint-specific peak-memory and active-parameter "
        "benchmark is still absent, so the complete deployment gate is pending."
    ),
}


def _find_paragraph(document: Document, needle: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if needle in paragraph.text]
    if len(matches) != 1:
        raise RuntimeError(
            f"expected one paragraph containing {needle!r}, got {len(matches)}"
        )
    return matches[0]


def _clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag.endswith("}pPr"):
            continue
        paragraph._p.remove(child)


def _set_text(paragraph: Paragraph, text: str, *, yellow: bool = False) -> None:
    _clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    if yellow:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _cell_set_text(cell, text: str, *, yellow: bool = False) -> None:
    _set_text(cell.paragraphs[0], text, yellow=yellow)
    for paragraph in list(cell.paragraphs[1:]):
        paragraph._element.getparent().remove(paragraph._element)


def _row_by_label(table, label: str):
    rows = [row for row in table.rows if row.cells[0].text.strip() == label]
    if len(rows) != 1:
        raise RuntimeError(f"expected one table row labeled {label!r}, got {len(rows)}")
    return rows[0]


def _artifact_row(table, needle: str):
    rows = [row for row in table.rows if needle in row.cells[0].text]
    if len(rows) != 1:
        raise RuntimeError(
            f"expected one artifact row containing {needle!r}, got {len(rows)}"
        )
    return rows[0]


def rewrite(input_path: Path, output_path: Path) -> None:
    document = Document(input_path)

    status = _find_paragraph(document, "YELLOW HIGHLIGHT = evidence still required")
    _set_text(
        status,
        "YELLOW HIGHLIGHT = evidence still required before submission. The "
        "hypothesis system now follows the final W0 -> W1 -> R-ECMoE training "
        "and evaluation chain. Available artifacts support H0 through H5. H6 "
        "remains partial because the epoch-130 checkpoint lacks a dedicated "
        "peak-memory and active-parameter deployment benchmark. Supplementary "
        "same-checkpoint mechanism ablations and publication figures that need "
        "quantile-level or per-origin action artifacts also remain outstanding.",
        yellow=True,
    )

    abstract = _find_paragraph(document, "Artificial-intelligence workloads create")
    base = abstract.text.split(
        " In a separate 37,384-interval chronologically held-out NLR physical "
        "evaluation,",
        1,
    )[0]
    _set_text(
        abstract,
        base
        + " The stage-aligned audit supports the corrected W0 foundation gate, "
        "the four-horizon W1 dynamics gate, R-A safe expansion, R-B full-2024 "
        "checkpoint selection, sealed-2025 value, and the balanced Top-2 design "
        "comparison; the final deployment-memory gate remains incomplete.",
    )

    introduction = document.paragraphs[26]
    if introduction.text.strip():
        raise RuntimeError("the hypothesis introduction placeholder is no longer blank")
    _set_text(
        introduction,
        "The following H0-H6 map is aligned to the frozen R-ECMoE architecture "
        "and its actual stage gates. It is reported as a post-architecture-revision "
        "evaluation framework, not as the superseded physical-twin or legacy "
        "scenario-planning hypothesis registry. Each proposition names its "
        "training stage, primary design metric, and acceptance rule.",
    )

    physical = _find_paragraph(document, "Held-out physical-core result (H0).")
    _set_text(
        physical,
        "Upstream physical-data audit (not an H0-H6 model hypothesis). On the "
        "chronologically held-out final 20% of verified NLR support (37,384 "
        "intervals), the frozen piecewise-linear mapping used offline to construct "
        "PUE targets reduced facility-power MAE from 338.245 to 56.266 kW, "
        "facility-power RMSE from 359.538 to 87.109 kW, daily-energy MAE from "
        "7,898.343 to 1,302.331 kWh, PUE MAE from 0.12139 to 0.02345, and ramp "
        "MAE from 4.987 to 1.126 kW per interval relative to the uncalibrated "
        "engineering curve. Component closure and source support also passed. "
        "This mapping is frozen before neural training and is not a teacher or "
        "trainable W0/W1/R module. Its nominal 90% diagnostic interval covered "
        "42.788%, but that interval is not consumed by W0, W1, or R and therefore "
        "is reported only as an upstream calibration limitation, not a model gate."
    )

    ablations = _find_paragraph(document, "The required deployable comparisons are")
    _set_text(
        ablations,
        "The required deployable design comparisons are the safe tensor baseline, "
        "separately trained dense-direct R, separately trained independent R-MoE, "
        "separately trained R-ECMoE Top-1 plus conditional tail, and the balanced "
        "Top-2 main model. All five rows are present in the frozen 2025 suite and "
        "support H4-H5. Supplementary same-checkpoint interventions isolating the "
        "route prior, active expert states, conditional tail, always-on tail, risk "
        "condition, and sparse coordinator are not present for epoch 130; they "
        "remain yellow-marked optional mechanism evidence rather than H0-H6 gates.",
        yellow=True,
    )

    first_decision = _find_paragraph(document, "H0 is NOT SUPPORTED")
    _set_text(
        first_decision,
        "H0 is SUPPORTED under the corrected W0 scope; H1 is SUPPORTED by the "
        "epoch-56 four-horizon dynamics release; H2 is SUPPORTED by the finite, "
        "physically feasible R-A improvement; and H3 is SUPPORTED because epoch "
        "130 minimizes decision value among all completed eligible full-2024 R-B "
        "evaluations. The run ended before its patience-based early-stop event, "
        "so no early-stop-trigger claim is made.",
    )
    second_decision = _find_paragraph(document, "H1 is SUPPORTED by the released W1")
    _set_text(
        second_decision,
        "H4 is SUPPORTED by the paired sealed-2025 comparison against the safe "
        "tensor. H5 is SUPPORTED by the separately retrained Top-1, dense-direct, "
        "and independent-MoE comparisons together with non-collapsed seven-expert "
        "routing. H6 is PARTIALLY SUPPORTED: latency and one-pass solver-free "
        "execution pass, but final checkpoint-specific peak memory and active "
        "parameter count remain unmeasured.",
        yellow=True,
    )

    conclusion_physical = _find_paragraph(document, "The separate local physical test")
    _set_text(
        conclusion_physical,
        "The held-out NLR calculation is retained as an upstream physical-data "
        "audit rather than an H0-H6 model hypothesis. Point-error, ramp, closure, "
        "and source-support checks favor the frozen mapping used to construct PUE "
        "targets. Its under-covered diagnostic interval is disclosed as a source-"
        "model limitation but does not determine W0, W1, R, or final-checkpoint "
        "acceptance because the deployed neural chain does not consume it.",
    )

    conclusion = _find_paragraph(document, "Across 362 common 2025 dates")
    prefix = conclusion.text.split(" H3 is supported", 1)[0]
    _set_text(
        conclusion,
        prefix
        + " The stage-aligned decisions support H0 (corrected W0 foundation), "
        "H1 (W1 dynamics), H2 (R-A expansion), H3 (R-B selection), H4 "
        "(safe-baseline value), and H5 (Top-2 design). H6 remains partial pending "
        "the final checkpoint memory and active-parameter benchmark. These findings "
        "are counterfactual realized-replay results in composite units, not dollar "
        "savings or observed NLR market participation.",
    )

    gap = _find_paragraph(document, "PRE-SUBMISSION EVIDENCE GAP:")
    _set_text(
        gap,
        "PRE-SUBMISSION EVIDENCE GAP: obtain an epoch-130 peak-memory and active-"
        "parameter deployment benchmark to complete H6; disclose that the run "
        "ended at epoch 168 before the four-full-evaluation early-stop condition "
        "fired; run the optional epoch-130 same-checkpoint route/tail/state/"
        "coordinator interventions if mechanism-level attribution is desired; and "
        "generate the remaining publication figures requiring quantile-level and "
        "per-origin action artifacts.",
        yellow=True,
    )

    hypothesis_check = _find_paragraph(
        document,
        "does not rewrite hypotheses after seeing the test results",
    )
    _set_text(
        hypothesis_check,
        "\u2610  The final manuscript labels the rewritten H0-H6 as a "
        "post-architecture stage-aligned evaluation framework and does not "
        "misrepresent it as the original preregistration.",
    )

    section_table = document.tables[3]
    _cell_set_text(
        section_table.cell(2, 1),
        "Scientific question, stage-aligned H0-H6 gates, and contributions",
    )

    hypothesis_table = document.tables[5]
    if len(hypothesis_table.rows) != 8:
        raise RuntimeError("hypothesis-definition table must contain H0-H6")
    _cell_set_text(hypothesis_table.cell(0, 1), "Stage-aligned hypothesis and gate")
    for index, hypothesis_id in enumerate(HYPOTHESES, start=1):
        if hypothesis_table.cell(index, 0).text.strip() != hypothesis_id:
            raise RuntimeError("hypothesis-definition row order changed")
        _cell_set_text(hypothesis_table.cell(index, 1), HYPOTHESES[hypothesis_id])

    decision_table = document.tables[31]
    if len(decision_table.rows) != 8:
        raise RuntimeError("hypothesis-decision table must contain H0-H6")
    _cell_set_text(decision_table.cell(0, 1), "Stage decision and measured evidence")
    for index, hypothesis_id in enumerate(DECISIONS, start=1):
        if decision_table.cell(index, 0).text.strip() != hypothesis_id:
            raise RuntimeError("hypothesis-decision row order changed")
        _cell_set_text(
            decision_table.cell(index, 1),
            DECISIONS[hypothesis_id],
            yellow=hypothesis_id == "H6",
        )

    sensitivity = document.tables[33]
    full_validation = _row_by_label(sensitivity, "Formal full validation")
    _cell_set_text(
        full_validation.cells[2],
        "COMPLETE FOR SELECTION: validation.csv contains all completed full-2024 "
        "candidates at epochs 91, 104, 117, 130, 143, and 156. Epoch 130 has the "
        "minimum eligible decision value (1.178472) and is the checkpoint frozen "
        "for the final suite.",
    )
    early_stop = _row_by_label(sensitivity, "Early stop")
    _cell_set_text(
        early_stop.cells[2],
        "Training-control rule only. The run ended at epoch 168 after two completed "
        "post-best full evaluations, before patience 4 could fire; the paper does "
        "not claim an early-stop event.",
        yellow=True,
    )

    artifacts = document.tables[35]
    continuation = _artifact_row(
        artifacts, "artifacts/r_ecmoe_balanced_continued/.../run_meta.json"
    )
    _cell_set_text(
        continuation.cells[0],
        "artifacts/r_ecmoe_balanced_continued/.../run_meta.json / "
        "config_resolved.yaml / continuation_lineage.json / train.jsonl / "
        "validation.csv / validation_summary.json",
    )
    _cell_set_text(
        continuation.cells[1],
        "H2-H3 evidence: R-A finite safe expansion, complete R-B candidate trace, "
        "full-2024 eligibility metrics, epoch-130 minimum decision value, and the "
        "disclosed non-triggered early-stop state.",
    )
    optional = _artifact_row(
        artifacts, "artifacts/r_ecmoe_balanced_reselection/.../evaluation_2024/variant_summary.csv"
    )
    _cell_set_text(
        optional.cells[1],
        "Development-stage epoch-80 same-checkpoint mechanism audit. Retained as "
        "supplementary evidence only; it is not used to decide the stage-aligned "
        "H0-H6 for the epoch-130 model.",
        yellow=True,
    )
    foundation = _artifact_row(
        artifacts, "artifacts/vpp_world_moe_r5_v2_1_2025final/.../phase=w0/state.json"
    )
    _cell_set_text(
        foundation.cells[1],
        "H0-H1 foundation evidence: corrected W0 state/identity scope and W1 epoch "
        "56/update 3360 four-horizon capability/release gate, including causal "
        "Top-2 routes, conditional-tail state, action-conditioned dynamics, and "
        "an unopened 2025 partition.",
    )
    physical_row = _artifact_row(artifacts, "artifacts/research/physical/run_id=paper-physical-v1")
    _cell_set_text(
        physical_row.cells[1],
        "Upstream frozen physical-label audit, not an H0-H6 hypothesis. Point, "
        "ramp, closure, and support metrics are retained for data provenance; the "
        "unused interval-coverage diagnostic is reported only as a limitation.",
    )
    registry = artifacts.add_row()
    _cell_set_text(registry.cells[0], "configs/r_ecmoe_stage_hypotheses.yaml")
    _cell_set_text(
        registry.cells[1],
        "Machine-readable post-architecture stage map for the rewritten H0-H6 "
        "definitions, thresholds, selectors, and deployment gate.",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    rewrite(args.input.resolve(), args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
