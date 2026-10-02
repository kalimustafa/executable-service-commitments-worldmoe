from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

from docx import Document


OLD_TEXT = (
    "Use one global random seed list fixed before experiments; report mean, "
    "standard deviation, and every failed seed."
)
NEW_TEXT = (
    "Use one global random seed list fixed before experiments. Under the "
    "2026-08-11 protocol amendment, every confirmatory training, replay, "
    "optimization, sensitivity, and paired-inference panel uses exactly "
    "[20260809, 20260819, 20260829, 20260908]. Seed 20260918 is excluded "
    "before its pending CfC-W trial and formal DART training; any earlier "
    "artifact containing it is retained only as "
    "legacy_pre_amendment_excluded_from_confirmatory_four_seed_panel. "
    "Report the four-seed mean, standard deviation, and every failed seed."
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def replace_paragraph_text(paragraph, old: str, new: str) -> bool:
    if paragraph.text == new:
        return False
    if paragraph.text != old:
        return False
    if not paragraph.runs:
        paragraph.add_run(new)
        return True
    paragraph.runs[0].text = new
    for run in paragraph.runs[1:]:
        run.text = ""
    return True


def main() -> None:
    repository = Path(__file__).resolve().parents[1]
    workspace = repository.parent
    source = workspace / (
        "Data_Center_VPP_English_Manuscript_and_Experimental_Protocol_"
        "NLR_NYISO_Revised.docx"
    )
    backup_dir = repository / "artifacts" / "manifests" / "specification_backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / (
        "Data_Center_VPP_English_Manuscript_and_Experimental_Protocol_"
        "NLR_NYISO_Revised.pre_four_seed_amendment_20260811.docx"
    )
    manifest_path = (
        repository / "artifacts" / "manifests" / "four_seed_docx_amendment.json"
    )

    before_sha = sha256(source)
    if not backup.exists():
        shutil.copy2(source, backup)

    document = Document(str(source))
    before_structure = {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "sections": len(document.sections),
        "inline_shapes": len(document.inline_shapes),
    }
    matches = [p for p in document.paragraphs if p.text in {OLD_TEXT, NEW_TEXT}]
    if len(matches) != 1:
        raise RuntimeError(f"expected one seed-protocol paragraph, found {len(matches)}")
    changed = replace_paragraph_text(matches[0], OLD_TEXT, NEW_TEXT)
    if changed:
        temporary = source.with_suffix(".four-seed.tmp.docx")
        document.save(str(temporary))
        temporary.replace(source)

    amended = Document(str(source))
    after_structure = {
        "paragraphs": len(amended.paragraphs),
        "tables": len(amended.tables),
        "sections": len(amended.sections),
        "inline_shapes": len(amended.inline_shapes),
    }
    if before_structure != after_structure:
        raise RuntimeError("DOCX structure changed outside the requested paragraph")
    if sum(p.text == NEW_TEXT for p in amended.paragraphs) != 1:
        raise RuntimeError("four-seed protocol paragraph was not persisted exactly once")

    payload = {
        "schema_version": "dcvpp.four_seed_docx_amendment/v1",
        "status": "COMPLETE",
        "source_path": source.as_posix(),
        "source_sha256_before": before_sha,
        "source_sha256_after": sha256(source),
        "backup_path": backup.as_posix(),
        "backup_sha256": sha256(backup),
        "four_seed_panel": [20260809, 20260819, 20260829, 20260908],
        "excluded_legacy_seed": 20260918,
        "changed": changed,
        "structure_before": before_structure,
        "structure_after": after_structure,
        "render_or_layout_qa_performed": False,
        "render_qa_omitted_by_user_instruction": True,
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
