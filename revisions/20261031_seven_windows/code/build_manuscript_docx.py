from pathlib import Path
from copy import deepcopy
import subprocess,json,re,hashlib,zipfile
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1];OLD=ROOT.parent/'EnergyEconomics_Revision_20261031'
OUT=ROOT/'deliverables';OUT.mkdir(exist_ok=True);WORK=ROOT/'work';WORK.mkdir(exist_ok=True)
PANDOC=r'C:\Users\Kali\anaconda3\Scripts\pandoc.exe'
mainref=OLD/'source_package/manuscript/EnergyEconomics_Contract_Value_20261031.docx'
suppref=OLD/'deliverables/EnergyEconomics_Supplementary_Revised_20261031.docx'
mainout=OUT/'EnergyEconomics_Seven_Window_Revision_20261031.docx'
suppout=OUT/'EnergyEconomics_Seven_Window_Supplement_20261031.docx'

def convert(md,out,ref):
    subprocess.run([PANDOC,str(md),'--from=markdown+tex_math_dollars','--to=docx','--standalone',
        '--reference-doc='+str(ref),'--resource-path='+str(ROOT/'manuscript'),'-o',str(out)],check=True)
def text(p):return ''.join(t.text or '' for t in p._p.xpath('.//w:t|.//m:t'))
def style(d,name):return next(s for s in d.styles if s.name==name)
def setup_table(t,widths=None,template=None):
    cols=len(t._tbl.tr_lst[0].tc_lst)
    if template:
        t._tbl.replace(t._tbl.tblPr,deepcopy(template._tbl.tblPr));t._tbl.replace(t._tbl.tblGrid,deepcopy(template._tbl.tblGrid))
        for r0,r1 in zip(template._tbl.tr_lst,t._tbl.tr_lst):
            for c0,c1 in zip(r0.tc_lst,r1.tc_lst):
                if c1.tcPr is not None:c1.remove(c1.tcPr)
                if c0.tcPr is not None:c1.insert(0,deepcopy(c0.tcPr))
    else:
        widths=widths or [6.5/cols]*cols;grid=t._tbl.tblGrid
        for e in list(grid):grid.remove(e)
        for w in widths:
            e=OxmlElement('w:gridCol');e.set(qn('w:w'),str(round(1440*w)));grid.append(e)
        for row in t.rows:
            for i,cell in enumerate(row.cells):cell.width=Inches(widths[i])
    t.autofit=False
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
        if i==0 and pr.find(qn('w:tblHeader')) is None:pr.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
                p.paragraph_format.keep_with_next=i<len(t.rows)-1
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9.5)

convert(ROOT/'manuscript/main.md',mainout,mainref)
d=Document(mainout);base=Document(mainref);d.paragraphs[0].style=style(d,'Title')
for i,t in enumerate(d.tables):
    setup_table(t,template=base.tables[i]) if i<4 else setup_table(t,[.75,1.15,1.35,1.95,1.3])
refs=False
for p in d.paragraphs:
    tx=text(p)
    if tx=='References':refs=True
    elif refs:
        p.paragraph_format.space_after=Pt(4)
        for r in p.runs:r.font.size=Pt(10)
    if tx.startswith(('Fig. ','Table ')):
        p.style=style(d,'Caption');p.paragraph_format.keep_together=True
        p.paragraph_format.keep_with_next=tx.startswith('Table ')
        p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(7)
        for r in p.runs:r.font.size=Pt(9.5)
    if p._p.xpath('.//w:drawing'):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(4)
    if p._p.xpath('.//m:oMathPara'):p.paragraph_format.keep_together=True
d.save(mainout)

sd=Document(suppref);oldtables=[t._tbl.xml for t in sd.tables]
captions={
1:'High-price import displacement and routine ramps. Dotted lines mark immediate execution; panels retain original historical trajectories.',
2:'Native-system financial differences and savings concentration. Paired intervals use seven-day blocks; removed-event totals are not rescheduled.',
3:'Historical service-window and nonpreemptive outcomes. The retained 30 h zero label is superseded by 27 overdue jobs in Table 15.',
4:'Monthly NYC tariff savings. Energy and demand components follow Tables 23–24 and 31.',
5:'Monthly savings across regional price portfolios. Values average three refinements; transfer uses frozen weights.',
6:'Detailed WorldMoE implementation. Dashed paths show training and solid paths execution; S23 defines the archived information boundary.',
7:'WorldMoE dispatch and independent certification. Routing is schematic; the original artwork is retained.',
8:'Original comparisons and WorldMoE ablations. Ranges and interventions follow Tables 9–10 and 28–29; historical certificate annotations are retained.',
9:'Historical opportunity and deadline comparisons. Old witness and ceiling labels are retained; current actual-admission certificates are in S31.',
10:'Frozen-weight ISO-NE transfer. Regional prices change while the facilities, workload and settlement conventions remain matched.',
11:'NYC tariff sensitivity. Bars show demand-charge changes and lines total saving; negative bars indicate additional charges.',
12:'Annual saving and paired physical-valuation gains. Bootstrap intervals and ceiling labels are historical; current opportunity bounds are in S31.'}
changed=[]
for p in sd.paragraphs:
    tx=text(p);m=re.match(r'Supplementary (?:Fig\.|Figure) (\d+)',tx)
    if m:
        n=int(m.group(1));p.text=f'Supplementary Fig. {n}. '+captions[n];changed.append(n)
        p.paragraph_format.keep_together=True
        for r in p.runs:r.font.size=Pt(9)
    if tx.startswith('Revision status as of 31 October 2026.'):
        p.text='Revision status as of 31 October 2026. The recovered service evidence and original ablation tables and images are retained. S31 adds seven actual-admission opportunity certificates; S32 adds new fixed-checkpoint policy evaluations, including the substantive 8 h service failure. All figure captions have been shortened. The earlier admission-compatible bounds remain valid robustness statements; the current main-text endpoints use the tighter actual-admission certificates.'
    if tx.startswith('S30.5 Remaining opportunity'):
        note=sd.add_paragraph('This subsection retains the earlier uniform-bound calculation. For the recovered actual admissions, S31 now supplies the tighter 24 h interval and the updated headroom and capture measures used in the main text.')
        p._p.addnext(note._p)
assert sorted(changed)==list(range(1,13))

addition=(ROOT/'certificates/S31_actual_admission_seven_window_certificates.md').read_text(encoding='utf-8')+'\n\n'+(ROOT/'S32_window_extension_methods.md').read_text(encoding='utf-8')
suppresults=ROOT/'manuscript/S32_results_appendix.md'
assert suppresults.exists();addition+='\n\n'+suppresults.read_text(encoding='utf-8')
addition=re.sub(r'\\tag\{([^}]+)\}',lambda m:r'\qquad\text{('+m.group(1)+')}',addition)
addmd=WORK/'supplement_additions.md';addmd.write_text(addition,encoding='utf-8')
adddoc=WORK/'supplement_additions.docx';convert(addmd,adddoc,suppref)
ad=Document(adddoc)
for el in ad.element.body:
    if el.tag!=qn('w:sectPr'):sd.element.body.insert(len(sd.element.body)-1,deepcopy(el))
for p in sd.paragraphs:
    if text(p).startswith(('S31 ','S32 Extension')):p.paragraph_format.page_break_before=True
    if p._p.xpath('.//m:oMathPara'):p.paragraph_format.keep_together=True
for t in sd.tables[len(oldtables):]:
    n=len(t._tbl.tr_lst[0].tc_lst)
    setup_table(t,[.9,1.85,1.85,1.9] if n==4 else ([1.1,2.6,2.8] if n==3 else None))
sd.save(suppout)
assert all(t._tbl.xml==old for t,old in zip(sd.tables,oldtables))
def media(p):
    with zipfile.ZipFile(p) as z:return {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
assert media(suppref)<=media(suppout)
manifest={'main':{'paragraphs':len(d.paragraphs),'tables':len(d.tables),'figures':len(d.inline_shapes),'native_math':len(d.element.xpath('.//m:oMath'))},
    'supplement':{'paragraphs':len(sd.paragraphs),'tables':len(sd.tables),'figures':len(sd.inline_shapes),'native_math':len(sd.element.xpath('.//m:oMath'))},
    'original_43_supplement_tables_preserved':True,'original_12_images_preserved':True,'shortened_supplement_captions':changed}
(WORK/'document_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8');print(json.dumps(manifest,indent=2))
